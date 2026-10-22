"""Шаблоны блоков «по умолчанию», комплектация ТКП, характеристики модулей, сброс пароля."""

import io
from unittest.mock import AsyncMock, patch

from openpyxl import load_workbook

from scr.dbase import crud_modules, crud_users
from scr.dbase.crud_tkp import kit_line
from scr.dbase.models import (
    AccuracyClass,
    Material,
    MaterialType,
    Module,
    ModuleItem,
    ModuleType,
    ModuleTypeDefault,
    ModuleTypeSlot,
    Request,
    User,
)
from scr.utils.excel import MONEY_FMT


async def _setup(s, sfx: str):
    """Раздел ЩО-типа с двумя строками (вторая скрытая), модули 100 и 50 ₽, старый модуль по умолчанию без шаблона."""
    user = User(name=f"Tpl{sfx}", email=f"tpl{sfx}@test.ru", hashed_password="x")
    s.add(user)
    await s.flush()
    req = Request(manager_id=user.id)
    mt = ModuleType(name=f"Щиты {sfx}", slug=f"tpl-{sfx}")
    m1, m2 = Module(name=f"ВВ-{sfx}"), Module(name=f"ТТ-{sfx}")
    mat1, mat2 = Material(name=f"Мат1-{sfx}", price=100), Material(name=f"Мат2-{sfx}", price=50)
    s.add_all([req, mt, m1, m2, mat1, mat2])
    await s.flush()
    breaker = ModuleTypeSlot(module_type_id=mt.id, name="Выключатель", sort_order=1, purpose="Вакуумный выключатель")
    hidden = ModuleTypeSlot(module_type_id=mt.id, name="Корпус", sort_order=2, hidden=True)
    s.add_all([breaker, hidden])
    await s.flush()
    s.add_all(
        [
            ModuleItem(module_id=m1.id, material_id=mat1.id, quantity=1),
            ModuleItem(module_id=m2.id, material_id=mat2.id, quantity=1),
            ModuleTypeDefault(module_type_id=mt.id, module_id=m1.id, quantity=1, slot_id=breaker.id),
        ]
    )
    await s.commit()
    return {"req": req.id, "mt": mt.id, "m1": m1.id, "m2": m2.id, "breaker": breaker.id, "hidden": hidden.id}


async def test_templates_and_block_from_template(client, db_session):
    ids = await _setup(db_session, "a")

    # Старые модули по умолчанию переходят в основной шаблон
    tpls = (await client.get(f"/reg/api/module-types/{ids['mt']}/templates")).json()
    assert [t["name"] for t in tpls] == ["Основной"]
    main_id = tpls[0]["id"]

    # Новый шаблон — копия основного, со своим назначением и модулем вне строк
    resp = await client.post(
        f"/reg/api/module-types/{ids['mt']}/templates", json={"name": "ЩО сх1", "copy_from": main_id}
    )
    tpl = resp.json()
    assert tpl["slots"][str(ids["breaker"])]["module_id"] == ids["m1"]
    await client.patch(f"/reg/api/module-type-templates/{tpl['id']}", json={"purpose": "Ввод"})
    resp = await client.put(
        f"/reg/api/module-type-slots/{ids['breaker']}/default",
        json={"module_id": ids["m2"], "quantity": 3, "template_id": tpl["id"]},
    )
    assert resp.status_code == 200
    await client.put(
        f"/reg/api/module-type-slots/{ids['hidden']}/default",
        json={"module_id": ids["m1"], "quantity": 1, "template_id": tpl["id"]},
    )
    tpl = (
        await client.post(
            f"/reg/api/module-type-templates/{tpl['id']}/items", json={"module_id": ids["m1"], "quantity": 2}
        )
    ).json()
    assert [(i["module_id"], i["quantity"]) for i in tpl["items"]] == [(ids["m1"], 2)]

    # Основной шаблон не изменился
    main = await crud_modules_default(db_session, main_id, ids["breaker"])
    assert (main.module_id, main.quantity) == (ids["m1"], 1)

    # Блок из шаблона: название, назначение, модули строк и добавленные; скрытая строка — в стоимости
    resp = await client.post(
        "/reg/api/request-calcs/",
        json={"request_id": ids["req"], "module_type_id": ids["mt"], "template_id": tpl["id"]},
    )
    block = resp.json()["block"]
    assert (block["name"], block["purpose"]) == ("ЩО сх1", "Ввод")
    assert block["slots"][0]["item"]["module_id"] == ids["m2"] and block["slots"][0]["item"]["quantity"] == 3
    assert block["hidden_total"] == 100.0
    assert [(i["module_id"], i["quantity"]) for i in block["items"]] == [(ids["m1"], 2)]

    # Блок без шаблона — из основного
    resp = await client.post("/reg/api/request-calcs/", json={"request_id": ids["req"], "module_type_id": ids["mt"]})
    assert resp.json()["block"]["name"] == "Основной"

    # Последний шаблон не удаляется
    assert (await client.delete(f"/reg/api/module-type-templates/{tpl['id']}")).status_code == 200
    assert (await client.delete(f"/reg/api/module-type-templates/{main_id}")).status_code == 400


async def crud_modules_default(s, template_id: int, slot_id: int) -> ModuleTypeDefault:
    from sqlalchemy import select

    stmt = select(ModuleTypeDefault).where(
        ModuleTypeDefault.template_id == template_id, ModuleTypeDefault.slot_id == slot_id
    )
    return (await s.execute(stmt)).scalar_one()


def test_kit_line():
    assert kit_line("Вакуумный выключатель", 1, "ESQ ВВ(F)-12-25/1250") == "Вакуумный выключатель: ESQ ВВ(F)-12-25/1250"
    assert kit_line("Трансформаторы тока", 3, "ТПЛ-10") == "Трансформаторы тока: 3хТПЛ-10"
    assert kit_line(None, 1, "Дуговая защита Лайм") == "1хДуговая защита Лайм"


async def test_tkp_kit_page_and_excel(client, db_session):
    ids = await _setup(db_session, "b")
    tpl = (await client.post(f"/reg/api/module-types/{ids['mt']}/templates", json={"name": "КСО-204м сх7.1ВВ"})).json()
    await client.put(
        f"/reg/api/module-type-slots/{ids['breaker']}/default",
        json={"module_id": ids["m1"], "quantity": 1, "template_id": tpl["id"]},
    )
    block = (
        await client.post(
            "/reg/api/request-calcs/",
            json={"request_id": ids["req"], "module_type_id": ids["mt"], "template_id": tpl["id"]},
        )
    ).json()["block"]
    await client.put(
        f"/reg/api/request-calcs/{block['id']}",
        json={"id": block["id"], "purpose": "Ввод", "quantity": 2, "final_price": 1000, "total_cost": 800},
    )

    resp = await client.get(f"/reg/api/requests/{ids['req']}/tkp-kit/export-excel")
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content)).active
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    assert rows[0][1] == "КСО-204м сх7.1ВВ-Ввод, \nВакуумный выключатель: ВВ-b"
    assert rows[0][2:5] == (2, 1000, "=D3*C3")
    assert ws.cell(3, 4).number_format == MONEY_FMT

    client.cookies.set("user_email", "tplb@test.ru")
    resp = await client.get(f"/reg/requests/{ids['req']}/calc/tkp")
    client.cookies.clear()
    assert "Комплектация" in resp.text and "КСО-204м сх7.1ВВ-Ввод" in resp.text
    assert "Вакуумный выключатель: ВВ-b" in resp.text


def test_money_format_separates_all_groups():
    # «# ##0.00» отделяло только последнюю тройку
    assert MONEY_FMT == "#,##0.00"


async def test_module_attrs_and_sort(db_session):
    s = db_session
    mtype = MaterialType(name="ТТсорт", color="#ff0000")
    s.add_all([mtype, AccuracyClass(name="0.5S-sort"), AccuracyClass(name="0.2-sort")])
    await s.flush()
    mats = [
        Material(name="ТТ А", price=1, type_id=mtype.id, accuracy_class="0.2-sort", voltage="10", vykat=True),
        Material(name="ТТ Б", price=1, type_id=mtype.id, accuracy_class="0.5S-sort", voltage="10"),
        Material(name="Болт", price=1, accuracy_class="5P", el_priv=True),
    ]
    mods = [
        Module(name="Сорт Я", nom_tok=400, material_type_id=mtype.id),
        Module(name="Сорт Б", nom_tok=400, material_type_id=mtype.id),
        Module(name="Сорт А", nom_tok=630, material_type_id=mtype.id),
    ]
    s.add_all(mats + mods)
    await s.flush()
    s.add_all(
        [
            ModuleItem(module_id=mods[0].id, material_id=mats[0].id, quantity=1),
            ModuleItem(module_id=mods[0].id, material_id=mats[2].id, quantity=1),
            ModuleItem(module_id=mods[1].id, material_id=mats[1].id, quantity=1),
            ModuleItem(module_id=mods[2].id, material_id=mats[0].id, quantity=1),
        ]
    )
    await s.commit()
    attrs = await crud_modules.get_module_attrs(s)
    # Свойства — по материалам типа модуля: болт другого типа не учитывается
    assert attrs[mods[0].id]["accuracy_class"] == "0.2-sort"
    assert attrs[mods[0].id]["vykat"] is True and attrs[mods[0].id]["el_priv"] is False
    ordered = await crud_modules.sort_modules(s, mods, attrs)
    # Ток, затем класс точности по справочнику (0.5S-sort раньше 0.2-sort), затем название
    assert [m.name for m in ordered] == ["Сорт Б", "Сорт Я", "Сорт А"]


async def test_material_type_color(client):
    t = (await client.post("/reg/api/material-types/", json={"name": "Цвет"})).json()
    resp = await client.patch("/reg/api/material-types/", json={"id": t["id"], "color": "#12abef"})
    assert resp.json()["color"] == "#12abef"
    resp = await client.patch("/reg/api/material-types/", json={"id": t["id"], "color": "red;x"})
    assert resp.status_code == 422
    resp = await client.patch("/reg/api/material-types/", json={"id": t["id"], "color": None})
    assert resp.json()["color"] is None


async def test_password_reset(client, db_session):
    user = User(name="Сброс", email="reset@test.ru", hashed_password=crud_users.hash_password("old-pass"))
    db_session.add(user)
    await db_session.commit()

    # Только для вошедших; письмо уходит со ссылкой на страницу сброса
    assert (await client.post(f"/reg/api/users/{user.id}/reset-password")).status_code == 401
    client.cookies.set("user_email", "reset@test.ru")
    with patch("scr.utils.email.send_password_reset", new=AsyncMock(return_value=None)) as sent:
        resp = await client.post(f"/reg/api/users/{user.id}/reset-password")
    client.cookies.clear()
    assert resp.status_code == 200
    link = sent.call_args.args[3]
    assert "/reg/reset-password?token=" in link
    token = link.split("token=")[1]

    assert "reset@test.ru" in (await client.get(f"/reg/reset-password?token={token}")).text
    resp = await client.post("/reg/reset-password", data={"token": token, "password": "new-pass", "password2": "nope"})
    assert "Пароли не совпадают" in resp.text
    resp = await client.post(
        "/reg/reset-password", data={"token": token, "password": "new-pass", "password2": "new-pass"}
    )
    assert "Пароль изменён" in resp.text
    resp = await client.post("/reg/api/auth/login", json={"email": "reset@test.ru", "password": "new-pass"})
    assert resp.status_code == 200
    # Ссылка одноразовая: после смены пароля недействительна
    assert "недействительна" in (await client.get(f"/reg/reset-password?token={token}")).text
    assert "недействительна" in (await client.get("/reg/reset-password?token=garbage")).text


async def test_pages_render(client, db_session):
    ids = await _setup(db_session, "c")
    mt = await db_session.get(ModuleType, ids["mt"])
    client.cookies.set("user_email", "tplc@test.ru")
    try:
        for url in ("/reg/modules", "/reg/requests", "/reg/material-types", "/reg/users", "/reg/contracts"):
            resp = await client.get(url)
            assert resp.status_code == 200, url
            assert "data-filter-table" in resp.text, url
        resp = await client.get(f"/reg/module-types-calc/{ids['mt']}/defaults")
        assert resp.status_code == 200 and "module-picker.js" in resp.text
        resp = await client.get(f"/reg/requests/{ids['req']}/calc/{mt.slug}")
        assert resp.status_code == 200 and 'id="moduleFilter"' in resp.text
        resp = await client.get(f"/reg/module-type-slots?type={mt.slug}")
        assert resp.status_code == 200 and "Назначение" in resp.text
    finally:
        client.cookies.clear()
