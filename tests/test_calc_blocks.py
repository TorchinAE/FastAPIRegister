import io

from openpyxl import Workbook, load_workbook

from scr.dbase.models import (
    Directors,
    Material,
    Module,
    ModuleItem,
    ModuleType,
    ModuleTypeDefault,
    ModuleTypeSlot,
    Organization,
    Positions,
    Request,
    User,
)


def _xlsx(rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return {"file": ("f.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


async def _setup(s, sfx: str = ""):
    """Заявка, материал 100 ₽ / 630 А, подмодуль (2 материала) и модуль (3 подмодуля) — 600 ₽."""
    user = User(name=f"Calc{sfx}", email=f"calc@test.ru{sfx}", hashed_password="x")
    pos = Positions(name=f"Calc position{sfx}")
    s.add_all([user, pos])
    await s.flush()
    director = Directors(name=f"Calc Director{sfx}", position_id=pos.id)
    s.add(director)
    await s.flush()
    org = Organization(name=f"Calc Org{sfx}", director_id=director.id, profitability=10)
    s.add(org)
    await s.flush()
    req = Request(manager_id=user.id, company_id=org.id)
    mat = Material(name=f"CalcMat{sfx}", price=100, nom_tok=630, code_1c=f"CALC-1C{sfx}")
    sub = Module(name=f"CalcSub{sfx}")
    top = Module(name=f"CalcTop{sfx}")
    mt = ModuleType(name=f"Calc KSO{sfx}", slug=f"calc-kso{sfx}")
    s.add_all([req, mat, sub, top, mt])
    await s.flush()
    slot = ModuleTypeSlot(module_type_id=mt.id, name="Шкаф", sort_order=1)
    s.add(slot)
    await s.flush()
    s.add_all(
        [
            ModuleItem(module_id=sub.id, material_id=mat.id, quantity=2),
            ModuleItem(module_id=top.id, sub_module_id=sub.id, quantity=3),
            ModuleTypeDefault(module_type_id=mt.id, module_id=top.id, quantity=1, slot_id=slot.id),
        ]
    )
    await s.commit()
    return {"req": req.id, "mat": mat.id, "sub": sub.id, "top": top.id, "mt": mt.id}


async def test_calc_blocks_prices_and_section_cost(client, db_session):
    ids = await _setup(db_session)

    # Новый блок: модули обязательных строк с ценой с учётом подмодулей, рентабельность организации
    resp = await client.post("/reg/api/request-calcs/", json={"request_id": ids["req"], "module_type_id": ids["mt"]})
    assert resp.status_code == 200
    block = resp.json()["block"]
    assert block["slots"][0]["item"]["module_price"] == 600.0
    assert block["profitability_percent"] == 10.0
    assert block["final_price"] == 660.0

    # Добавленный модуль сразу возвращается с ценой
    calc_id = block["id"]
    resp = await client.post(
        f"/reg/api/request-calcs/{calc_id}/items", json={"calc_id": calc_id, "module_id": ids["sub"], "quantity": 1}
    )
    assert resp.json()["module_price"] == 200.0

    # Второй блок с количеством: стоимость раздела = сумма(общая стоимость × кол-во)
    resp = await client.post("/reg/api/request-calcs/", json={"request_id": ids["req"], "module_type_id": ids["mt"]})
    second = resp.json()["block"]["id"]
    resp = await client.put(
        f"/reg/api/request-calcs/{second}", json={"id": second, "final_price": 1000, "total_cost": 900, "quantity": 3}
    )
    assert resp.status_code == 200
    resp = await client.delete(f"/reg/api/request-calcs/{second}")
    assert resp.json()["section_cost"] == 660.0

    # Комплектация: модули ×кол-во блоков, материалы с раскрытием подмодулей, итоги с прибылью
    resp = await client.get(f"/reg/api/requests/{ids['req']}/kit/export-excel")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Итоги", "Комплектация", "Материалы"]
    modules = {r[2]: r for r in wb["Комплектация"].iter_rows(min_row=3, values_only=True)}
    assert modules["CalcTop"][3] == 1 and modules["CalcSub"][3] == 1
    materials = {r[3]: r for r in wb["Материалы"].iter_rows(min_row=3, values_only=True)}
    assert materials["CalcMat"][6] == 8  # 3×2 из CalcTop + 2 из CalcSub
    totals = {r[0]: r for r in wb["Итоги"].iter_rows(min_row=3, values_only=True)}
    assert totals["Итого"][8] == 660.0 and totals["Итого"][7] == 660.0 - totals["Итого"][5]

    client.cookies.set("user_email", "calc@test.ru")
    resp = await client.get(f"/reg/requests/{ids['req']}/calc")
    client.cookies.clear()
    assert "Выгрузить комплектацию" in resp.text and "Средняя рентабельность" in resp.text


async def test_module_import_excel_and_nom_tok(client, db_session):
    ids = await _setup_module_only(db_session)
    files = _xlsx(
        [
            ["Название", "Цена за ед.", "Кол-во", "Тип материала", "Код 1С"],
            ["другое имя", 1, 7, None, "IMP-1C"],
            ["Импорт новый", 55.5, 2, None, None],
        ]
    )
    resp = await client.post(f"/reg/api/modules/{ids['mod']}/import-excel", files=files)
    assert resp.status_code == 200
    data = resp.json()
    assert (data["added"], data["created_materials"]) == (2, 1)

    resp = await client.get(f"/reg/api/modules/{ids['mod']}")
    items = {i["material"]["name"]: i for i in resp.json()["items"]}
    assert items["ImpMat"]["quantity"] == 7  # найден по коду 1С
    new_mat = await client.get(f"/reg/api/materials/{items['Импорт новый']['material_id']}")
    assert new_mat.json()["type_name"] == "Производство"

    # Ном ток модуля — максимальный из материалов, можно задать вручную и вернуть авто
    modules = {m["id"]: m for m in (await client.get("/reg/api/modules/all")).json()}
    assert modules[ids["mod"]]["nom_tok"] == 1600
    await client.patch("/reg/api/modules/", json={"id": ids["mod"], "nom_tok": 400})
    modules = {m["id"]: m for m in (await client.get("/reg/api/modules/all")).json()}
    assert modules[ids["mod"]]["nom_tok"] == 400
    assert modules[ids["mod"]]["nom_tok_manual"] is True
    await client.patch("/reg/api/modules/", json={"id": ids["mod"], "nom_tok_manual": False})
    modules = {m["id"]: m for m in (await client.get("/reg/api/modules/all")).json()}
    assert modules[ids["mod"]]["nom_tok"] == 1600


async def _setup_module_only(s):
    mat = Material(name="ImpMat", price=10, nom_tok=1600, code_1c="IMP-1C")
    mod = Module(name="ImpModule")
    s.add_all([mat, mod])
    await s.commit()
    return {"mat": mat.id, "mod": mod.id}


async def test_defaults_import_creates_material_and_module(client, db_session):
    mt = ModuleType(name="Def KSO", slug="def-kso")
    db_session.add_all([mt, Module(name="DefExisting")])
    await db_session.flush()
    db_session.add_all(
        [
            ModuleTypeSlot(module_type_id=mt.id, name="Ввод", sort_order=1),
            ModuleTypeSlot(module_type_id=mt.id, name="Шкаф", sort_order=2),
        ]
    )
    await db_session.commit()
    ids = {"mt": mt.id}
    files = _xlsx(
        [
            ["Строка", "ID", "Модуль", "Кол-во", "Цена за ед."],
            ["Ввод", None, "DefExisting", 2, None],
            ["Шкаф", None, "Новый шкаф", 3, 250],
            [None, None, "DefExisting", 1, None],
        ]
    )
    resp = await client.post(f"/reg/api/module-types/{ids['mt']}/defaults/import-excel", files=files)
    assert resp.status_code == 200
    data = resp.json()
    # Модули по умолчанию — только обязательные строки: строка Excel без «Строки» пропускается
    assert (data["added"], data["created_modules"], data["created_materials"]) == (2, 1, 1)
    assert data["skipped"] == ["DefExisting: нет обязательной строки «»"]

    defaults = {d["module_name"]: d for d in (await client.get(f"/reg/api/module-types/{ids['mt']}/defaults")).json()}
    assert defaults["Новый шкаф"]["quantity"] == 3
    assert defaults["Новый шкаф"]["module_price"] == 250.0
    mats = (await client.get("/reg/api/materials/?search=Новый шкаф")).json()["items"]
    assert [(m["name"], m["type_name"]) for m in mats] == [("Новый шкаф", "Производство")]


async def test_kso_split_qty_and_calc_excel(client, db_session):
    ids = await _setup(db_session, "2")
    mt = await db_session.get(ModuleType, ids["mt"])
    mt.slug = "kso"
    await db_session.commit()

    # КСО: новый блок — 1 шт. на 1СШ; общее количество = 1СШ + 2СШ
    resp = await client.post("/reg/api/request-calcs/", json={"request_id": ids["req"], "module_type_id": ids["mt"]})
    block = resp.json()["block"]
    assert (block["qty_1sh"], block["qty_2sh"], block["quantity"]) == (1, 0, 1)
    resp = await client.put(
        f"/reg/api/request-calcs/{block['id']}", json={"id": block["id"], "qty_1sh": 2, "qty_2sh": 3}
    )
    assert resp.json()["quantity"] == 5

    # Сводная по материалам: блок — кол-во на один блок, итого — с учётом кол-ва блоков
    resp = await client.get(f"/reg/api/requests/{ids['req']}/calc/kso/export-excel")
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content)).active
    assert "Calc Org2" in ws.cell(1, 1).value and "Calc KSO2" in ws.cell(1, 1).value
    assert ws.cell(2, 5).value == "Calc KSO2 1 (1СШ: 2, 2СШ: 3)"
    rows = {r[1]: r for r in ws.iter_rows(min_row=4, values_only=True)}
    assert rows["CalcMat2"][:3] == (ids["mat"], "CalcMat2", 100)
    assert rows["CalcMat2"][4:8] == (6, 600, 30, 3000)
    assert rows["Итого материалы"][5] == 600 and rows["Итого материалы"][7] == 3000
    assert rows["Конечная цена за блок"][5] == 660
    assert rows["Количество блоков"][4] == 5
    assert rows["Конечная цена × кол-во"][7] == 3300

    # Модули по умолчанию: ID — первая колонка, внизу общая стоимость
    resp = await client.get(f"/reg/api/module-types/{ids['mt']}/defaults/export-excel")
    ws = load_workbook(io.BytesIO(resp.content)).active
    assert [c.value for c in ws[2]] == ["ID", "Строка", "Модуль", "Кол-во", "Цена за ед.", "Сумма"]
    assert ws.cell(3, 1).value == ids["top"]
    assert ws.cell(4, 2).value == "Общая стоимость" and ws.cell(4, 6).value == "=SUM(F3:F3)"
