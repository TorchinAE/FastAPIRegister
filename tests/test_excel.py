import io

import pytest
from openpyxl import Workbook


@pytest.mark.asyncio
async def test_export_materials_excel(client):
    await client.post("/reg/api/materials/", json={"name": "ExcelTest", "price": 42})
    response = await client.get("/reg/api/materials/export-excel")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    wb = load_workbook_from_bytes(response.content)
    ws = wb.active
    assert ws.cell(1, 1).value == "Материалы"
    assert ws.cell(2, 1).value == "ID"
    assert ws.cell(2, 2).value == "Название"


@pytest.mark.asyncio
async def test_import_materials_excel(client):
    wb = Workbook()
    ws = wb.active
    ws.append(["Название", "Цена", "Код 1С", "Код агент", "URL агент"])
    ws.append(["ИмпортМат1", 100, "C1", "A1", "http://test.ru"])
    ws.append(["ИмпортМат2", 200, "C2", "A2", ""])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    files = {"file": ("materials.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    response = await client.post("/reg/api/materials/import-excel", files=files)
    assert response.status_code == 200
    data = response.json()
    assert data["imported"] == 2
    resp = await client.get("/reg/api/materials/?search=ИмпортМат1")
    assert resp.json()["total"] >= 1


@pytest.mark.asyncio
async def test_export_modules_excel(client):
    await client.post("/reg/api/modules/", json={"name": "ExcelModule"})
    response = await client.get("/reg/api/modules/export-excel")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]


@pytest.mark.asyncio
async def test_import_modules_excel(client):
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Модули"
    ws1.append(["ID", "Название"])
    ws1.append([1, "ИмпортМодуль1"])
    ws2 = wb.create_sheet("Состав")
    ws2.append(["Модуль", "Тип компонента", "Название компонента", "Количество"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    files = {"file": ("modules.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    response = await client.post("/reg/api/modules/import-excel", files=files)
    assert response.status_code == 200
    assert response.json()["imported"] >= 1


def load_workbook_from_bytes(content: bytes):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(content))
