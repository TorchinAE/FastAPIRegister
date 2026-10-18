import io
import re

from openpyxl import load_workbook

from scr.dbase.models import Directors, Organization, Positions, Request, User


async def _setup(s, sfx: str):
    user = User(name=f"Pay{sfx}", email=f"pay{sfx}@test.ru", hashed_password="x")
    pos = Positions(name=f"Pay position{sfx}")
    s.add_all([user, pos])
    await s.flush()
    director = Directors(name=f"Pay Director{sfx}", position_id=pos.id)
    s.add(director)
    await s.flush()
    org = Organization(name=f"Pay Org{sfx}", director_id=director.id)
    s.add(org)
    await s.flush()
    req = Request(manager_id=user.id, company_id=org.id, cost=100000, factory_order_num=f"26-{sfx}")
    s.add(req)
    await s.flush()
    req.tkp_num = f"{req.id}-ив"
    await s.commit()
    return user.email, req.id


async def test_payment_items_saved_after_opening_page(client, db_session):
    """Порядок оплаты: строки, созданные при открытии страницы ТКП, сохраняются после ввода."""
    email, req_id = await _setup(db_session, "page")
    client.cookies.set("user_email", email)
    try:
        resp = await client.get(f"/reg/requests/{req_id}")
        assert resp.status_code == 200
        rows = re.findall(r'class="payment-row" data-id="(\d+)" data-type="([^"]+)"', resp.text)
        ids = {ptype: int(pid) for pid, ptype in rows}
        assert set(ids) == {"предоплата", "доплата-1", "доплата-2", "доплата-3", "отсрочка"}

        payload = {
            "id": ids["предоплата"],
            "amount": 30000,
            "percent": 30,
            "due_date": "2026-11-01",
            "paid_amount": 10000,
            "paid_date": "2026-10-20",
        }
        resp = await client.put(f"/reg/api/payments/{ids['предоплата']}", json=payload)
        assert resp.status_code == 200

        # После повторного открытия страницы введённое значение на месте, строки не пересоздаются
        resp = await client.get(f"/reg/requests/{req_id}")
        assert f'data-id="{ids["предоплата"]}" data-type="предоплата"' in resp.text
        assert 'value="30000.00"' in resp.text
    finally:
        client.cookies.clear()

    items = {p["payment_type"]: p for p in (await client.get(f"/reg/api/payments/by-request/{req_id}")).json()}
    assert len(items) == 5
    assert items["предоплата"]["amount"] == 30000
    assert items["предоплата"]["paid_amount"] == 10000
    assert items["предоплата"]["due_date"].startswith("2026-11-01")


async def test_payment_registry_excel(client, db_session):
    _, req_id = await _setup(db_session, "reg")
    items = {p["payment_type"]: p["id"] for p in (await client.get(f"/reg/api/payments/by-request/{req_id}")).json()}
    await client.put(f"/reg/api/payments/{items['предоплата']}", json={"id": 0, "amount": 50000, "paid_amount": 50000})
    await client.put(f"/reg/api/payments/{items['доплата-1']}", json={"id": 0, "amount": 50000, "paid_amount": 20000})

    resp = await client.post("/reg/api/payments/registry-excel", json={"ids": [req_id], "filter_text": "ТКП: reg"})
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["Реестр оплат", "По ТКП", "По видам платежей"]
    ws = wb["Реестр оплат"]
    assert "фильтр: ТКП: reg" in ws["A1"].value
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    assert [r[4] for r in rows[:2]] == ["предоплата", "доплата-1"]
    assert rows[0][1] == "26-reg"  # номер заказа
    subtotal = rows[2]
    assert subtotal[0].startswith("Итого по ТКП") and subtotal[6] == 100000 and subtotal[8] == 70000
    assert subtotal[10] == 30000
    per_tkp = list(wb["По ТКП"].iter_rows(min_row=3, values_only=True))
    assert per_tkp[0][1] == "26-reg" and per_tkp[0][6] == 30000
