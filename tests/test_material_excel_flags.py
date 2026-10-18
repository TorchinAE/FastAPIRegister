import io

from openpyxl import Workbook, load_workbook


def _xlsx(rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return {"file": ("m.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


async def _material(client, name: str) -> dict:
    items = (await client.get(f"/reg/api/materials/?search={name}&per_page=100")).json()["items"]
    return next(m for m in items if m["name"] == name)


async def _flags(db_session, mat_id: int) -> tuple:
    from scr.dbase.models import Material

    mat = await db_session.get(Material, mat_id)
    await db_session.refresh(mat)
    return mat.stats, mat.vtych, mat.vykat, mat.ruchn, mat.el_priv


async def test_import_bool_flags_from_exported_layout(client, db_session):
    """Признаки стац/втыч/выкат/ручн/эл.прив: 1 или + → да, 0 или - → нет, в том числе у существующих материалов.

    Файл в формате выгрузки: строка-заголовок «Материалы» над шапкой.
    """
    await client.post("/reg/api/materials/", json={"name": "ФлагиСущ", "price": 10, "nom_tok": 630})
    existing = await _material(client, "ФлагиСущ")
    header = ["ID", "Тип", "Название", "Цена", "Ном Ток", "стац", "втыч", "выкат", "ручн", "эл.прив"]
    files = _xlsx(
        [
            ["Материалы"],
            header,
            [existing["id"], None, "ФлагиСущ", 10, 630, "-", 1, "+", 0, "1"],
            [None, None, "ФлагиНовый", 5, 0, "+", "-", 0, "+", None],
        ]
    )
    resp = await client.post("/reg/api/materials/import-excel-confirm", files=files)
    assert resp.status_code == 200
    assert resp.json()["imported"] == 2

    assert await _flags(db_session, existing["id"]) == (False, True, True, False, True)
    # Без ном. тока значения из файла всё равно учитываются, пустая ячейка — «нет»
    new = await _material(client, "ФлагиНовый")
    assert await _flags(db_session, new["id"]) == (True, False, False, True, False)


async def test_export_materials_filtered_by_ids(client):
    ids = []
    for name in ("ВыгрузкаА", "ВыгрузкаБ", "ВыгрузкаВ"):
        ids.append((await client.post("/reg/api/materials/", json={"name": name, "price": 1})).json()["id"])
    resp = await client.post("/reg/api/materials/export-excel", json={"ids": [ids[2], ids[0]]})
    assert resp.status_code == 200
    ws = load_workbook(io.BytesIO(resp.content)).active
    names = [r[2] for r in ws.iter_rows(min_row=3, values_only=True)]
    assert names == ["ВыгрузкаВ", "ВыгрузкаА"]


async def test_bulk_modules_for_filtered_materials(client):
    ids = []
    for name in ("МодМассА", "МодМассБ"):
        ids.append((await client.post("/reg/api/materials/", json={"name": name, "price": 3})).json()["id"])
    resp = await client.post("/reg/api/materials/modules-bulk?create=true", json={"ids": ids})
    assert resp.status_code == 200
    assert sorted(resp.json()["changed"]) == sorted(ids)
    modules = (await client.get("/reg/api/modules/?search=МодМасс&per_page=100")).json()["items"]
    assert {m["name"] for m in modules} >= {"МодМассА", "МодМассБ"}

    # Повторная установка ничего не меняет; снятие удаляет модули
    assert (await client.post("/reg/api/materials/modules-bulk?create=true", json={"ids": ids})).json()["changed"] == []
    resp = await client.post("/reg/api/materials/modules-bulk?create=false", json={"ids": ids})
    assert sorted(resp.json()["changed"]) == sorted(ids)
    modules = (await client.get("/reg/api/modules/?search=МодМасс&per_page=100")).json()["items"]
    assert not [m for m in modules if m["name"] in ("МодМассА", "МодМассБ")]
