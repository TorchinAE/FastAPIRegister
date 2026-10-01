import io

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from scr.dbase import crud_material_types, crud_materials
from scr.dbase.database import db_helper
from scr.dbase.schemas.schemas import (
    MaterialCreateSchema,
    MaterialResponseSchema,
    MaterialTypeCreateSchema,
    MaterialUpdateSchema,
    PaginatedResponse,
)

mat_router = APIRouter(prefix="/api/materials", tags=["Materials"])


def _mat_response(mat) -> MaterialResponseSchema:
    return MaterialResponseSchema(
        id=mat.id,
        name=mat.name,
        price=float(mat.price),
        code_1c=mat.code_1c,
        code_agent=mat.code_agent,
        url_agent=mat.url_agent,
        type_id=mat.type_id,
        type_name=mat.type.name if mat.type else None,
        created_by=mat.created_by,
        date=mat.date,
    )


@mat_router.post("/", response_model=MaterialResponseSchema)
async def add_material(
    data: MaterialCreateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    new_mat = await crud_materials.add_material(session=session, in_mat=data)
    await session.commit()
    await session.refresh(new_mat, ["type"])
    return _mat_response(new_mat)


@mat_router.get("/", response_model=PaginatedResponse)
async def read_materials(
    search: str | None = Query(None),
    page: int = Query(1, ge=1),
    per_page: int = Query(20, ge=1, le=100),
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    items, total = await crud_materials.get_materials(session=session, search=search, page=page, per_page=per_page)
    return PaginatedResponse(
        items=[_mat_response(i) for i in items],
        total=total,
        page=page,
        per_page=per_page,
        pages=(total + per_page - 1) // per_page,
    )


@mat_router.get("/all", response_model=list[MaterialResponseSchema])
async def read_all_materials(
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    items = await crud_materials.get_all_materials(session=session)
    return [_mat_response(i) for i in items]


@mat_router.get("/export-excel")
async def export_materials_excel(
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    from datetime import datetime as dt

    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    items = await crud_materials.get_all_materials(session=session)
    wb = Workbook()
    ws = wb.active
    ws.title = "Материалы"

    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Title row
    headers = ["ID", "Тип", "Название", "Цена", "Дата", "Ном Ток", "стац", "втыч", "выкат", "ручн", "эл.прив", "Код 1С", "Код агент", "URL агент"]
    num_cols = len(headers)
    ws.append(["Материалы"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
    ws.cell(1, 1).font = Font(bold=True, size=14)
    ws.cell(1, 1).alignment = Alignment(horizontal="center")

    ws.append(headers)
    for col_idx in range(1, num_cols + 1):
        cell = ws.cell(2, col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    for item in items:
        ws.append(
            [
                item.id,
                item.type.name if item.type else "",
                item.name,
                float(item.price),
                item.date.strftime("%d.%m.%Y") if item.date else "",
                item.nom_tok,
                "+" if item.stats else "",
                "+" if item.vtych else "",
                "+" if item.vykat else "",
                "+" if item.ruchn else "",
                "+" if item.el_priv else "",
                item.code_1c or "",
                item.code_agent or "",
                item.url_agent or "",
            ]
        )
        row_num = ws.max_row
        ws.cell(row_num, 4).number_format = "# ##0.00"
        for col_idx in range(1, num_cols + 1):
            ws.cell(row_num, col_idx).border = thin_border

    # Auto-fit columns
    from openpyxl.utils import get_column_letter

    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value is not None:
                cell_len = len(str(cell.value))
                max_len = max(max_len, cell_len)
        ws.column_dimensions[col_letter].width = min(max_len + 3, 60)

    date_str = dt.now().strftime("%Y%m%d")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"materials_{date_str}.xlsx"
    from urllib.parse import quote

    encoded = quote(filename)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"},
    )


@mat_router.post("/import-excel")
async def import_materials_excel(
    file: UploadFile = File(...),
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    content = await file.read()
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    # Phase 1: scan for conflicts (match by id + price)
    conflicts = []
    new_rows = []
    for row in rows:
        if not row or not row[0]:
            continue
        name = str(row[0]).strip()
        if not name:
            continue
        price = float(row[1]) if len(row) > 1 and row[1] else 0
        row_id = int(row[6]) if len(row) > 6 and row[6] else None

        existing = None
        if row_id:
            existing = await crud_materials.get_material_by_id(session, row_id)
        if existing and float(existing.price) == price:
            conflicts.append(
                {
                    "id": existing.id,
                    "name": name,
                    "old_name": existing.name,
                    "price": price,
                    "old_price": float(existing.price),
                }
            )
        new_rows.append((row, row_id, existing))

    if conflicts:
        return {"status": "confirm", "conflicts": conflicts, "total": len(new_rows)}

    # Phase 2: no conflicts — import all
    imported = await _do_import(session, rows)
    await session.commit()
    return {"status": "ok", "message": f"Импортировано: {imported}", "imported": imported}


@mat_router.post("/import-excel-confirm")
async def import_materials_excel_confirm(
    file: UploadFile = File(...),
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    """Import with overwrite confirmation — user already confirmed."""
    content = await file.read()
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    imported = await _do_import(session, rows)
    await session.commit()
    return {"status": "ok", "message": f"Импортировано: {imported}", "imported": imported}


async def _do_import(session: AsyncSession, rows: list) -> int:
    from datetime import datetime as dt

    # Cache all material types for auto-creation
    types_cache, _ = await crud_material_types.get_all_types(session, per_page=10000)
    types_map = {t.name: t.id for t in types_cache}

    imported = 0
    for row in rows:
        if not row or not row[0]:
            continue

        # New format (14 cols): id, type, name, price, date, nom_tok, stats, vtych, vykat, ruchn, el_priv, code_1c, code_agent, url_agent
        # Old format (7 cols): name, price, code_1c, code_agent, url_agent, type_name, id
        if len(row) >= 14:
            row_id = int(row[0]) if row[0] else None
            type_name = str(row[1]).strip() if row[1] else None
            name = str(row[2]).strip() if row[2] else None
            if not name:
                continue
            price = float(row[3]) if row[3] else 0
            date_val = None
            if row[4]:
                try:
                    if isinstance(row[4], dt):
                        date_val = row[4].date()
                    else:
                        date_val = dt.strptime(str(row[4]).strip(), "%d.%m.%Y").date()
                except (ValueError, AttributeError):
                    pass
            nom_tok = int(row[5]) if row[5] else 0
            stats = bool(row[6]) if row[6] is not None else True
            vtych = bool(row[7]) if row[7] is not None else False
            vykat = bool(row[8]) if row[8] is not None else False
            ruchn = bool(row[9]) if row[9] is not None else True
            el_priv = bool(row[10]) if row[10] is not None else False
            code_1c = str(row[11]).strip() if len(row) > 11 and row[11] else None
            code_agent = str(row[12]).strip() if len(row) > 12 and row[12] else None
            url_agent = str(row[13]).strip() if len(row) > 13 and row[13] else None
        else:
            name = str(row[0]).strip()
            if not name:
                continue
            price = float(row[1]) if len(row) > 1 and row[1] else 0
            code_1c = str(row[2]).strip() if len(row) > 2 and row[2] else None
            code_agent = str(row[3]).strip() if len(row) > 3 and row[3] else None
            url_agent = str(row[4]).strip() if len(row) > 4 and row[4] else None
            type_name = str(row[5]).strip() if len(row) > 5 and row[5] else None
            row_id = int(row[6]) if len(row) > 6 and row[6] else None
            date_val = None
            nom_tok = 0
            stats = True
            vtych = False
            vykat = False
            ruchn = True
            el_priv = False

        type_id = None
        if type_name:
            if type_name not in types_map:
                new_type = await crud_material_types.add_type(session, MaterialTypeCreateSchema(name=type_name))
                types_map[type_name] = new_type.id
            type_id = types_map[type_name]

        schema = MaterialCreateSchema(
            name=name, price=price, code_1c=code_1c, code_agent=code_agent, url_agent=url_agent,
            type_id=type_id, nom_tok=nom_tok, stats=stats, vtych=vtych, vykat=vykat,
            ruchn=ruchn, el_priv=el_priv, date=date_val,
        )

        if row_id:
            existing = await crud_materials.get_material_by_id(session, row_id)
            if existing:
                for field, value in schema.model_dump(exclude_unset=True).items():
                    if hasattr(existing, field):
                        setattr(existing, field, value)
                imported += 1
                continue

        await crud_materials.add_material(session, schema)
        imported += 1
    return imported


@mat_router.get("/{mat_id}", response_model=MaterialResponseSchema)
async def read_material(mat_id: int, session: AsyncSession = Depends(db_helper.session_dependency)):
    mat = await crud_materials.get_material_by_id(session=session, mat_id=mat_id)
    if not mat:
        raise HTTPException(status_code=404, detail="Материал не найден")
    return _mat_response(mat)


@mat_router.patch("/", response_model=MaterialResponseSchema)
async def update_material(
    data: MaterialUpdateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    result = await crud_materials.update_material(session=session, upd_mat=data)
    if not result:
        raise HTTPException(status_code=404, detail="Материал не найден")
    await session.commit()
    await session.refresh(result, ["type"])
    return _mat_response(result)


@mat_router.delete("/{mat_id}")
async def delete_material(mat_id: int, session: AsyncSession = Depends(db_helper.session_dependency)):
    result = await crud_materials.delete_material(session=session, mat_id=mat_id)
    if not result:
        raise HTTPException(status_code=404, detail="Материал не найден")
    await session.commit()
    return {"message": "Материал удалён"}
