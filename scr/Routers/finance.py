import io
from datetime import UTC, datetime

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from scr.dbase import crud_finance
from scr.dbase.database import db_helper
from scr.dbase.schemas.schemas import (
    InvoiceItemCreateSchema,
    InvoiceItemResponseSchema,
    InvoiceItemUpdateSchema,
    PaginatedResponse,
    PayrollItemCreateSchema,
    PayrollItemResponseSchema,
    PayrollItemUpdateSchema,
)

fin_router = APIRouter(prefix="/api/finance", tags=["Finance"])


# ==================== Накладные ====================


@fin_router.get("/invoices", response_model=PaginatedResponse)
async def read_invoice_items(
    page: int = Query(1, ge=1),
    per_page: int = Query(20, ge=1, le=100),
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    items, total = await crud_finance.get_invoice_items(session, page=page, per_page=per_page)
    return PaginatedResponse(
        items=[_invoice_response(i, session) for i in items],
        total=total,
        page=page,
        per_page=per_page,
        pages=(total + per_page - 1) // per_page,
    )


@fin_router.get("/invoices/all", response_model=list[InvoiceItemResponseSchema])
async def read_all_invoice_items(session: AsyncSession = Depends(db_helper.session_dependency)):
    items = await crud_finance.get_all_invoice_items(session)
    return [_invoice_response(i, session) for i in items]


@fin_router.post("/invoices", response_model=InvoiceItemResponseSchema)
async def add_invoice_item(
    data: InvoiceItemCreateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    item = await crud_finance.add_invoice_item(session, name=data.name, cost=data.cost)
    await session.commit()
    return _invoice_response(item, session)


@fin_router.patch("/invoices", response_model=InvoiceItemResponseSchema)
async def update_invoice_item(
    data: InvoiceItemUpdateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    item = await crud_finance.update_invoice_item(session, item_id=data.id, name=data.name, cost=data.cost)
    if not item:
        raise HTTPException(status_code=404, detail="Запись не найдена")
    await session.commit()
    return _invoice_response(item, session)


@fin_router.delete("/invoices/{item_id}")
async def delete_invoice_item(
    item_id: int,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    result = await crud_finance.delete_invoice_item(session, item_id)
    if not result:
        raise HTTPException(status_code=404, detail="Запись не найдена")
    await session.commit()
    return {"message": "Удалено"}


@fin_router.get("/invoices/export-excel")
async def export_invoices_excel(session: AsyncSession = Depends(db_helper.session_dependency)):
    items = await crud_finance.get_all_invoice_items(session)
    return _export_excel(items, "Накладные", _invoice_response_items)


@fin_router.post("/invoices/import-excel")
async def import_invoices_excel(
    file: UploadFile = File(...),
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    imported = await _import_excel(session, file, "invoice")
    await session.commit()
    return {"status": "ok", "message": f"Импортировано: {imported}", "imported": imported}


# ==================== ФОТ ====================


@fin_router.get("/payroll", response_model=PaginatedResponse)
async def read_payroll_items(
    page: int = Query(1, ge=1),
    per_page: int = Query(20, ge=1, le=100),
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    items, total = await crud_finance.get_payroll_items(session, page=page, per_page=per_page)
    return PaginatedResponse(
        items=[_payroll_response(i, session) for i in items],
        total=total,
        page=page,
        per_page=per_page,
        pages=(total + per_page - 1) // per_page,
    )


@fin_router.get("/payroll/all", response_model=list[PayrollItemResponseSchema])
async def read_all_payroll_items(session: AsyncSession = Depends(db_helper.session_dependency)):
    items = await crud_finance.get_all_payroll_items(session)
    return [_payroll_response(i, session) for i in items]


@fin_router.post("/payroll", response_model=PayrollItemResponseSchema)
async def add_payroll_item(
    data: PayrollItemCreateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    item = await crud_finance.add_payroll_item(session, name=data.name, cost=data.cost)
    await session.commit()
    return _payroll_response(item, session)


@fin_router.patch("/payroll", response_model=PayrollItemResponseSchema)
async def update_payroll_item(
    data: PayrollItemUpdateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    item = await crud_finance.update_payroll_item(session, item_id=data.id, name=data.name, cost=data.cost)
    if not item:
        raise HTTPException(status_code=404, detail="Запись не найдена")
    await session.commit()
    return _payroll_response(item, session)


@fin_router.delete("/payroll/{item_id}")
async def delete_payroll_item(
    item_id: int,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    result = await crud_finance.delete_payroll_item(session, item_id)
    if not result:
        raise HTTPException(status_code=404, detail="Запись не найдена")
    await session.commit()
    return {"message": "Удалено"}


@fin_router.get("/payroll/export-excel")
async def export_payroll_excel(session: AsyncSession = Depends(db_helper.session_dependency)):
    items = await crud_finance.get_all_payroll_items(session)
    return _export_excel(items, "ФОТ", _payroll_response_items)


@fin_router.post("/payroll/import-excel")
async def import_payroll_excel(
    file: UploadFile = File(...),
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    imported = await _import_excel(session, file, "payroll")
    await session.commit()
    return {"status": "ok", "message": f"Импортировано: {imported}", "imported": imported}


# ==================== Helpers ====================


def _invoice_response(item, session) -> dict:
    return {
        "id": item.id,
        "name": item.name,
        "cost": float(item.cost),
        "date_modified": item.date_modified,
        "changed_by_id": item.changed_by_id,
        "changed_by_name": None,
    }


def _payroll_response(item, session) -> dict:
    return {
        "id": item.id,
        "name": item.name,
        "cost": float(item.cost),
        "date_modified": item.date_modified,
        "changed_by_id": item.changed_by_id,
        "changed_by_name": None,
    }


def _invoice_response_items(item):
    return {"name": item.name, "cost": float(item.cost)}


def _payroll_response_items(item):
    return {"name": item.name, "cost": float(item.cost)}


def _export_excel(items, title, item_mapper):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = title

    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    headers = ["ID", "Название", "Стоимость"]
    num_cols = len(headers)
    ws.append(headers)
    for col_idx in range(1, num_cols + 1):
        cell = ws.cell(1, col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border

    for item in items:
        ws.append([item.id, item.name, float(item.cost)])
        row_num = ws.max_row
        ws.cell(row_num, 3).number_format = "# ##0.00"
        for col_idx in range(1, num_cols + 1):
            ws.cell(row_num, col_idx).border = thin_border

    from openpyxl.utils import get_column_letter

    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 3, 60)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{title}.xlsx"},
    )


async def _import_excel(session, file, kind: str) -> int:
    content = await file.read()
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    imported = 0
    for row in rows:
        if not row:
            continue
        # Skip header
        first = str(row[0]).strip().lower() if row[0] else ""
        if first in ("id", "название", "имя"):
            continue

        name = str(row[0]).strip() if row[0] else None
        if not name:
            continue

        cost = 0
        if len(row) > 1 and row[1] is not None:
            try:
                cost = float(row[1])
            except (ValueError, TypeError):
                continue

        # If row has ID (col 0 is number), try to update existing
        item_id = None
        try:
            item_id = int(row[0])
        except (ValueError, TypeError):
            pass

        if item_id and kind == "invoice":
            existing = await crud_finance.get_invoice_item_by_id(session, item_id)
            if existing:
                await crud_finance.update_invoice_item(session, item_id, name=name, cost=cost)
                imported += 1
                continue
        elif item_id and kind == "payroll":
            existing = await crud_finance.get_payroll_item_by_id(session, item_id)
            if existing:
                await crud_finance.update_payroll_item(session, item_id, name=name, cost=cost)
                imported += 1
                continue

        if kind == "invoice":
            await crud_finance.add_invoice_item(session, name=name, cost=cost)
        else:
            await crud_finance.add_payroll_item(session, name=name, cost=cost)
        imported += 1

    return imported
