import io

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from scr.dbase import crud_module_slots, crud_module_templates, crud_module_types, crud_modules
from scr.dbase.database import db_helper
from scr.dbase.models import ModuleTypeDefault
from scr.dbase.schemas.schemas import (
    ModuleTypeCreateSchema,
    ModuleTypeDefaultCreateSchema,
    ModuleTypeDefaultResponseSchema,
    ModuleTypeResponseSchema,
    ModuleTypeSlotCreateSchema,
    ModuleTypeSlotUpdateSchema,
    ModuleTypeTemplateCreateSchema,
    ModuleTypeTemplateUpdateSchema,
    RequestCalcCreateSchema,
    RequestCalcItemCreateSchema,
    RequestCalcItemResponseSchema,
    RequestCalcItemUpdateSchema,
    RequestCalcResponseSchema,
    RequestCalcUpdateSchema,
    SlotModuleSchema,
    TemplateItemCreateSchema,
)

mt_calc_router = APIRouter(tags=["ModuleTypes"])


# ── Module Types ──


@mt_calc_router.get("/api/module-types/", response_model=list[ModuleTypeResponseSchema])
async def list_module_types(session: AsyncSession = Depends(db_helper.session_dependency)):
    types = await crud_module_types.get_module_types(session)
    return [
        ModuleTypeResponseSchema(
            id=t.id,
            name=t.name,
            slug=t.slug,
            fot_item_id=t.fot_item_id,
            overhead_item_id=t.overhead_item_id,
            fot_item_name=t.fot_item.name if t.fot_item else None,
            fot_item_cost=float(t.fot_item.cost) if t.fot_item else 0,
            overhead_item_name=t.overhead_item.name if t.overhead_item else None,
            overhead_item_cost=float(t.overhead_item.cost) if t.overhead_item else 0,
            created_by=t.created_by,
        )
        for t in types
    ]


@mt_calc_router.post("/api/module-types/", response_model=ModuleTypeResponseSchema)
async def create_module_type(
    data: ModuleTypeCreateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    obj = await crud_module_types.add_module_type(session, data)
    await session.commit()
    return obj


@mt_calc_router.delete("/api/module-types/{type_id}")
async def delete_module_type(
    type_id: int,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    result = await crud_module_types.delete_module_type(session, type_id)
    if not result:
        raise HTTPException(status_code=404, detail="Тип модуля не найден")
    await session.commit()
    return {"ok": True}


@mt_calc_router.put("/api/module-types/{type_id}", response_model=ModuleTypeResponseSchema)
async def update_module_type(
    type_id: int,
    data: dict,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    mt = await crud_module_types.get_module_type_by_id(session, type_id)
    if not mt:
        raise HTTPException(status_code=404, detail="Тип модуля не найден")
    if "fot_item_id" in data:
        mt.fot_item_id = int(data["fot_item_id"]) if data["fot_item_id"] else None
    if "overhead_item_id" in data:
        mt.overhead_item_id = int(data["overhead_item_id"]) if data["overhead_item_id"] else None
    await session.flush()
    await session.commit()
    mt = await crud_module_types.get_module_type_by_id(session, type_id)
    return ModuleTypeResponseSchema(
        id=mt.id,
        name=mt.name,
        slug=mt.slug,
        fot_item_id=mt.fot_item_id,
        overhead_item_id=mt.overhead_item_id,
        fot_item_name=mt.fot_item.name if mt.fot_item else None,
        fot_item_cost=float(mt.fot_item.cost) if mt.fot_item else 0,
        overhead_item_name=mt.overhead_item.name if mt.overhead_item else None,
        overhead_item_cost=float(mt.overhead_item.cost) if mt.overhead_item else 0,
        created_by=mt.created_by,
    )


# ── Module Type Defaults ──


@mt_calc_router.get("/api/module-types/{type_id}/defaults", response_model=list[ModuleTypeDefaultResponseSchema])
async def list_defaults(
    type_id: int,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    defaults = await crud_module_types.get_defaults(session, type_id)
    prices = await crud_modules.get_module_prices(session)
    return [
        ModuleTypeDefaultResponseSchema(
            id=d.id,
            module_type_id=d.module_type_id,
            module_id=d.module_id,
            quantity=d.quantity,
            module_name=d.module.name if d.module else None,
            module_price=prices.get(d.module_id, 0.0),
            slot_id=d.slot_id,
        )
        for d in defaults
    ]


@mt_calc_router.post("/api/module-types/{type_id}/defaults", response_model=ModuleTypeDefaultResponseSchema)
async def create_default(
    type_id: int,
    data: ModuleTypeDefaultCreateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    data.module_type_id = type_id
    obj = await crud_module_types.add_default(session, data)
    await session.commit()
    return ModuleTypeDefaultResponseSchema(
        id=obj.id,
        module_type_id=obj.module_type_id,
        module_id=obj.module_id,
        quantity=obj.quantity,
    )


_DEFAULTS_ALIASES = {
    "строка": "slot",
    "обязательная строка": "slot",
    "id": "id",
    "id модуля": "id",
    "модуль": "name",
    "название": "name",
    "наименование": "name",
    "кол-во": "qty",
    "количество": "qty",
    "код 1с": "code_1c",
    "код 1c": "code_1c",
    "тип материала": "mat_type",
    "цена за ед.": "price",
    "цена за ед": "price",
    "цена": "price",
}


@mt_calc_router.get("/api/module-types/{type_id}/defaults/export-excel")
async def export_defaults_excel(
    type_id: int,
    template_id: int | None = Query(None),
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    """Модули шаблона (не задан — основного) по обязательным строкам."""
    from datetime import datetime as dt

    from openpyxl.styles import Font

    from scr.utils.excel import MONEY_FMT, auto_fit_columns, border_row, safe_filename, write_header, xlsx_response

    mt = await crud_module_types.get_module_type_by_id(session, type_id)
    if not mt:
        raise HTTPException(status_code=404, detail="Тип модуля не найден")
    tpl = await crud_module_templates.resolve_template(session, type_id, template_id)
    slots = await crud_module_slots.get_slots(session, type_id)
    prices = await crud_modules.get_module_prices(session)
    by_slot = {d.slot_id: d for d in tpl.defaults if d.slot_id is not None}

    wb = Workbook()
    ws = wb.active
    ws.title = "Модули по умолчанию"
    headers = ["ID", "Строка", "Модуль", "Кол-во", "Цена за ед.", "Сумма"]
    ws.append([f"{mt.name} — модули по умолчанию, шаблон «{tpl.name}»"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(1, 1).font = Font(bold=True, size=14)
    write_header(ws, headers)
    first_row = ws.max_row + 1
    # Обязательные строки (включая скрытые) по порядку
    rows = [(s.name, by_slot.get(s.id)) for s in slots]
    for slot_name, d in rows:
        if d is None:
            ws.append([None, slot_name])
        else:
            values = [d.module_id, slot_name, d.module.name if d.module else "", d.quantity]
            ws.append([*values, prices.get(d.module_id, 0.0)])
        row_num = ws.max_row
        if d is not None:
            ws.cell(row_num, 6).value = f"=D{row_num}*E{row_num}"
        ws.cell(row_num, 5).number_format = MONEY_FMT
        ws.cell(row_num, 6).number_format = MONEY_FMT
        border_row(ws, row_num, len(headers))
    # Общая стоимость — в колонке «Строка»: при импорте строка без ID и модуля пропускается
    last_row = ws.max_row
    ws.append([None, "Общая стоимость", None, None, None, f"=SUM(F{first_row}:F{last_row})" if rows else 0])
    total_row = ws.max_row
    ws.cell(total_row, 6).number_format = MONEY_FMT
    for col in range(1, len(headers) + 1):
        ws.cell(total_row, col).font = Font(bold=True)
    border_row(ws, total_row, len(headers))
    auto_fit_columns(ws)
    name = f"defaults_{safe_filename(mt.slug)}_{safe_filename(tpl.name)}_{dt.now().strftime('%Y%m%d')}.xlsx"
    return xlsx_response(wb, name)


@mt_calc_router.post("/api/module-types/{type_id}/defaults/import-excel")
async def import_defaults_excel(
    type_id: int,
    file: UploadFile = File(...),
    template_id: int | None = Query(None),
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    """Загрузка в шаблон (не задан — в основной). Модуль ищется по ID, затем по названию.
    Уже добавленный модуль — количество заменяется.

    Колонка «Строка» с названием обязательной строки — модуль ставится в эту строку.
    Строки без обязательной строки пропускаются: модули по умолчанию — только обязательные строки.

    Модуль не найден — ищется (или создаётся) материал с этим названием (Код 1С, Цена за ед., Тип материала),
    из него создаётся модуль, который и добавляется в модули по умолчанию.
    """
    from scr.dbase.models import Module, ModuleItem
    from scr.Routers.modules import material_finder
    from scr.utils.excel import cell, data_rows, find_header, to_int

    mt = await crud_module_types.get_module_type_by_id(session, type_id)
    if not mt:
        raise HTTPException(status_code=404, detail="Тип модуля не найден")
    try:
        wb = load_workbook(io.BytesIO(await file.read()), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Не удалось прочитать файл Excel") from exc
    rows = list(wb.active.iter_rows(values_only=True))
    header_idx, col_map = find_header(rows, _DEFAULTS_ALIASES, required="name")
    if header_idx < 0:
        raise HTTPException(status_code=400, detail="Не найдена строка заголовка с колонкой «Модуль»")

    slots = {s.name.strip().lower(): s for s in await crud_module_slots.get_slots(session, type_id)}
    tpl = await crud_module_templates.resolve_template(session, type_id, template_id)
    find_or_create = await material_finder(session)
    added = updated = created_modules = created_materials = 0
    skipped: list[str] = []
    for row in data_rows(rows, header_idx):
        name = cell(row, col_map, "name")
        name = str(name).strip() if name is not None else None
        row_id = to_int(cell(row, col_map, "id"))
        if not name and not row_id:
            continue
        mod = await crud_modules.get_module_by_id(session, row_id) if row_id else None
        if not mod and name:
            mod = await crud_modules.get_module_by_name(session, name)
        if not mod and name:
            code_1c = cell(row, col_map, "code_1c")
            code_1c = str(code_1c).strip() if code_1c is not None else None
            mat, created = await find_or_create(
                None, code_1c, name, cell(row, col_map, "price"), cell(row, col_map, "mat_type")
            )
            created_materials += created
            mod = Module(name=name, material_type_id=mat.type_id)
            session.add(mod)
            await session.flush()
            session.add(ModuleItem(module_id=mod.id, material_id=mat.id, quantity=1))
            created_modules += 1
        if not mod:
            skipped.append(name or f"ID {row_id}")
            continue
        qty = to_int(cell(row, col_map, "qty"), 1)
        if qty is None or qty < 1:
            skipped.append(f"{mod.name}: количество {qty}")
            continue
        slot_name = cell(row, col_map, "slot")
        slot = slots.get(str(slot_name).strip().lower()) if slot_name else None
        if not slot:
            skipped.append(f"{mod.name}: нет обязательной строки «{slot_name or ''}»")
            continue
        default_stmt = select(ModuleTypeDefault.id).where(
            ModuleTypeDefault.slot_id == slot.id, ModuleTypeDefault.template_id == tpl.id
        )
        had_default = (await session.execute(default_stmt)).first() is not None
        try:
            await crud_module_slots.set_slot_default(session, slot, mod.id, qty, tpl.id)
        except ValueError as exc:
            skipped.append(f"{mod.name}: {exc}")
            continue
        if had_default:
            updated += 1
        else:
            added += 1
    if created_modules:
        await session.flush()
        await crud_modules.recalc_nom_tok(session)
    await session.commit()
    message = f"Добавлено: {added}, обновлено: {updated}"
    if created_modules:
        message += f", новых модулей: {created_modules}, новых материалов: {created_materials}"
    if skipped:
        message += f", не найдено: {', '.join(skipped)}"
    return {
        "message": message,
        "added": added,
        "updated": updated,
        "created_modules": created_modules,
        "created_materials": created_materials,
        "skipped": skipped,
    }


@mt_calc_router.patch("/api/module-types/defaults/{default_id}")
async def update_default(
    default_id: int,
    data: dict,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    obj = await session.get(ModuleTypeDefault, default_id)
    if not obj:
        raise HTTPException(status_code=404, detail="Запись не найдена")
    if "quantity" in data:
        obj.quantity = int(data["quantity"])
    await session.flush()
    await session.commit()
    return {"ok": True}


@mt_calc_router.delete("/api/module-types/defaults/{default_id}")
async def delete_default(
    default_id: int,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    result = await crud_module_types.delete_default(session, default_id)
    if not result:
        raise HTTPException(status_code=404, detail="Запись не найдена")
    await session.commit()
    return {"ok": True}


# ── Обязательные строки ──


async def _slot_or_404(session: AsyncSession, slot_id: int):
    slot = await crud_module_slots.get_slot_by_id(session, slot_id)
    if not slot:
        raise HTTPException(status_code=404, detail="Строка не найдена")
    return slot


@mt_calc_router.get("/api/module-types/{type_id}/slots")
async def list_slots(type_id: int, session: AsyncSession = Depends(db_helper.session_dependency)):
    return [crud_module_slots.slot_json(s) for s in await crud_module_slots.get_slots(session, type_id)]


@mt_calc_router.post("/api/module-types/{type_id}/slots")
async def create_slot(
    type_id: int,
    data: ModuleTypeSlotCreateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    if not await crud_module_types.get_module_type_by_id(session, type_id):
        raise HTTPException(status_code=404, detail="Тип модуля не найден")
    data.module_type_id = type_id
    slot = await crud_module_slots.add_slot(session, data)
    await session.commit()
    return crud_module_slots.slot_json(slot)


@mt_calc_router.patch("/api/module-type-slots/{slot_id}")
async def update_slot(
    slot_id: int,
    data: ModuleTypeSlotUpdateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    if not await crud_module_slots.update_slot(session, slot_id, data):
        raise HTTPException(status_code=404, detail="Строка не найдена")
    await session.commit()
    return crud_module_slots.slot_json(await _slot_or_404(session, slot_id))


@mt_calc_router.post("/api/module-type-slots/{slot_id}/move")
async def move_slot(slot_id: int, data: dict, session: AsyncSession = Depends(db_helper.session_dependency)):
    if not await crud_module_slots.move_slot(session, slot_id, int(data.get("direction", 0))):
        raise HTTPException(status_code=404, detail="Строка не найдена")
    await session.commit()
    return {"ok": True}


@mt_calc_router.delete("/api/module-type-slots/{slot_id}")
async def delete_slot(slot_id: int, session: AsyncSession = Depends(db_helper.session_dependency)):
    if not await crud_module_slots.delete_slot(session, slot_id):
        raise HTTPException(status_code=404, detail="Строка не найдена")
    await session.commit()
    return {"ok": True}


@mt_calc_router.put("/api/module-type-slots/{slot_id}/default")
async def set_slot_default(
    slot_id: int,
    data: SlotModuleSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    """Модуль обязательной строки в шаблоне «по умолчанию» (template_id не задан — в основном)."""
    slot = await _slot_or_404(session, slot_id)
    try:
        obj = await crud_module_slots.set_slot_default(session, slot, data.module_id, data.quantity, data.template_id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    await session.commit()
    if not obj:
        return {"item": None}
    prices = await crud_modules.get_module_prices(session)
    return {"item": {"id": obj.id, "module_id": obj.module_id, "quantity": obj.quantity,
                     "module_price": prices.get(obj.module_id, 0.0)}}  # fmt: skip


# ── Шаблоны блоков («по умолчанию») ──


async def _template_or_404(session: AsyncSession, template_id: int):
    tpl = await crud_module_templates.get_template(session, template_id)
    if not tpl:
        raise HTTPException(status_code=404, detail="Шаблон не найден")
    return tpl


async def _template_json(session: AsyncSession, template_id: int) -> dict:
    tpl = await _template_or_404(session, template_id)
    slots = await crud_module_slots.get_slots(session, tpl.module_type_id)
    prices = await crud_modules.get_module_prices(session)
    return crud_module_templates.template_json(tpl, slots, prices)


@mt_calc_router.get("/api/module-types/{type_id}/templates")
async def list_templates(type_id: int, session: AsyncSession = Depends(db_helper.session_dependency)):
    """Шаблоны раздела (id, название, назначение) — для выбора при добавлении блока."""
    templates = await crud_module_templates.get_templates(session, type_id)
    await session.commit()
    return [{"id": t.id, "name": t.name, "purpose": t.purpose or ""} for t in templates]


@mt_calc_router.post("/api/module-types/{type_id}/templates")
async def create_template(
    type_id: int,
    data: ModuleTypeTemplateCreateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    if not await crud_module_types.get_module_type_by_id(session, type_id):
        raise HTTPException(status_code=404, detail="Тип модуля не найден")
    tpl = await crud_module_templates.add_template(session, type_id, data.name, data.purpose, data.copy_from)
    await session.commit()
    return await _template_json(session, tpl.id)


@mt_calc_router.patch("/api/module-type-templates/{template_id}")
async def update_template(
    template_id: int,
    data: ModuleTypeTemplateUpdateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    tpl = await crud_module_templates.update_template(
        session, template_id, data.name, data.purpose, set_purpose="purpose" in data.model_fields_set
    )
    if not tpl:
        raise HTTPException(status_code=404, detail="Шаблон не найден")
    await session.commit()
    return {"id": tpl.id, "name": tpl.name, "purpose": tpl.purpose or ""}


@mt_calc_router.delete("/api/module-type-templates/{template_id}")
async def delete_template(template_id: int, session: AsyncSession = Depends(db_helper.session_dependency)):
    try:
        tpl = await crud_module_templates.delete_template(session, template_id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not tpl:
        raise HTTPException(status_code=404, detail="Шаблон не найден")
    await session.commit()
    return {"ok": True}


@mt_calc_router.post("/api/module-type-templates/{template_id}/items")
async def add_template_item(
    template_id: int,
    data: TemplateItemCreateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    """Модуль шаблона вне обязательных строк."""
    tpl = await _template_or_404(session, template_id)
    if not await crud_modules.get_module_by_id(session, data.module_id):
        raise HTTPException(status_code=404, detail="Модуль не найден")
    await crud_module_templates.add_template_item(session, tpl, data.module_id, data.quantity)
    await session.commit()
    return await _template_json(session, template_id)


@mt_calc_router.put("/api/request-calcs/{calc_id}/slots/{slot_id}")
async def set_calc_slot(
    calc_id: int,
    slot_id: int,
    data: SlotModuleSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    """Модуль обязательной строки в блоке расчёта. Скрытые строки задаются только в «по умолчанию»."""
    slot = await _slot_or_404(session, slot_id)
    calc = await crud_module_types.get_calc_by_id(session, calc_id)
    if not calc or calc.module_type_id != slot.module_type_id:
        raise HTTPException(status_code=404, detail="Расчёт не найден")
    if slot.hidden:
        raise HTTPException(status_code=400, detail="Скрытая строка задаётся в модулях по умолчанию")
    try:
        obj = await crud_module_slots.set_calc_slot(session, calc_id, slot, data.module_id, data.quantity)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    await session.commit()
    if not obj:
        return {"item": None}
    await session.refresh(obj, attribute_names=["module"])
    prices = await crud_modules.get_module_prices(session)
    return {
        "item": {
            "id": obj.id,
            "module_id": obj.module_id,
            "module_name": obj.module.name if obj.module else "?",
            "quantity": obj.quantity,
            "module_price": prices.get(obj.module_id, 0.0),
        }
    }


# ── RequestCalcs ──


@mt_calc_router.post("/api/request-calcs/")
async def create_calc(
    data: RequestCalcCreateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    """Новый блок расчёта со значениями по умолчанию. Возвращает данные блока для страницы."""
    if not await crud_module_types.get_module_type_by_id(session, data.module_type_id):
        raise HTTPException(status_code=404, detail="Тип модуля не найден")
    calc = await crud_module_types.create_calc(session, data.request_id, data.module_type_id, data.template_id)
    section_cost = await crud_module_types.recalc_section_cost(session, data.request_id, data.module_type_id)
    await session.commit()
    prices = await crud_modules.get_module_prices(session)
    default_profit = await crud_module_types.default_profitability(session, data.request_id)
    slots = await crud_module_slots.get_slots(session, data.module_type_id)
    return {
        "block": crud_module_types.calc_block_data(calc, prices, default_profit, slots),
        "section_cost": section_cost,
    }


@mt_calc_router.delete("/api/request-calcs/{calc_id}")
async def delete_calc(
    calc_id: int,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    calc = await crud_module_types.delete_calc(session, calc_id)
    if not calc:
        raise HTTPException(status_code=404, detail="Расчёт не найден")
    section_cost = await crud_module_types.recalc_section_cost(session, calc.request_id, calc.module_type_id)
    await session.commit()
    return {"ok": True, "section_cost": section_cost}


@mt_calc_router.get("/api/request-calcs/{calc_id}", response_model=RequestCalcResponseSchema)
async def get_calc(
    calc_id: int,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    calc = await crud_module_types.get_calc_by_id(session, calc_id)
    if not calc:
        raise HTTPException(status_code=404, detail="Расчёт не найден")
    return calc


@mt_calc_router.put("/api/request-calcs/{calc_id}", response_model=RequestCalcResponseSchema)
async def update_calc(
    calc_id: int,
    data: RequestCalcUpdateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    data.id = calc_id
    result = await crud_module_types.update_calc(session, data)
    if not result:
        raise HTTPException(status_code=404, detail="Расчёт не найден")
    await crud_module_types.recalc_section_cost(session, result.request_id, result.module_type_id)
    await session.commit()
    return result


@mt_calc_router.post("/api/request-calcs/{calc_id}/items", response_model=RequestCalcItemResponseSchema)
async def add_calc_item(
    calc_id: int,
    data: RequestCalcItemCreateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    data.calc_id = calc_id
    obj = await crud_module_types.add_calc_item(session, data)
    await session.commit()
    await session.refresh(obj, attribute_names=["module"])
    prices = await crud_modules.get_module_prices(session)
    return RequestCalcItemResponseSchema(
        id=obj.id,
        calc_id=obj.calc_id,
        module_id=obj.module_id,
        quantity=obj.quantity,
        module_name=obj.module.name if obj.module else None,
        module_price=prices.get(obj.module_id, 0.0),
    )


@mt_calc_router.patch("/api/request-calcs/items/{item_id}", response_model=RequestCalcItemResponseSchema)
async def update_calc_item(
    item_id: int,
    data: RequestCalcItemUpdateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    data.id = item_id
    result = await crud_module_types.update_calc_item(session, data)
    if not result:
        raise HTTPException(status_code=404, detail="Строка не найдена")
    await session.commit()
    await session.refresh(result, attribute_names=["module"])
    prices = await crud_modules.get_module_prices(session)
    return RequestCalcItemResponseSchema(
        id=result.id,
        calc_id=result.calc_id,
        module_id=result.module_id,
        quantity=result.quantity,
        module_name=result.module.name if result.module else None,
        module_price=prices.get(result.module_id, 0.0),
    )


@mt_calc_router.delete("/api/request-calcs/items/{item_id}")
async def delete_calc_item(
    item_id: int,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    result = await crud_module_types.delete_calc_item(session, item_id)
    if not result:
        raise HTTPException(status_code=404, detail="Строка не найдена")
    await session.commit()
    return {"ok": True}


@mt_calc_router.get("/api/requests/{req_id}/kit/export-excel")
async def export_kit_excel(req_id: int, session: AsyncSession = Depends(db_helper.session_dependency)):
    """Комплектация ТКП: итоги по разделам (ФОТ, накладные, рентабельность, прибыль), модули и материалы."""
    from datetime import datetime as dt

    from openpyxl.styles import Font

    from scr.dbase import crud_requests
    from scr.dbase.crud_tkp import build_kit, build_tkp_summary
    from scr.utils.excel import MONEY_FMT, auto_fit_columns, border_row, write_header, xlsx_response

    req = await crud_requests.get_request_by_id(session, req_id)
    if not req:
        raise HTTPException(status_code=404, detail="ТКП не найдена")
    if await crud_module_slots.sync_hidden_slots(session, req_id):
        await session.commit()
    summary = await build_tkp_summary(session, req)
    kit = await build_kit(session, req)
    title = f"ТКП #{req.id} {req.tkp_num or ''}".strip()
    bold = Font(bold=True)

    def sheet_title(ws, text: str, cols: int) -> None:
        ws.append([text])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
        ws.cell(1, 1).font = Font(bold=True, size=14)

    def add_row(ws, values: list, money_cols: tuple, cols: int, is_total: bool = False) -> None:
        ws.append(values)
        row_num = ws.max_row
        for col in money_cols:
            ws.cell(row_num, col).number_format = MONEY_FMT
        border_row(ws, row_num, cols)
        if is_total:
            for col in range(1, cols + 1):
                ws.cell(row_num, col).font = bold

    wb = Workbook()

    # Итоги по разделам
    ws = wb.active
    ws.title = "Итоги"
    headers = [
        "Раздел", "Кол-во", "Комплектация", "ФОТ", "Накладные", "Себестоимость",
        "Рентабельность, %", "Прибыль", "Итого без НДС",
    ]  # fmt: skip
    money = (3, 4, 5, 6, 8, 9)
    sheet_title(ws, f"{title} — итоги", len(headers))
    write_header(ws, headers)
    for r in summary["rows"]:
        values = [r["name"], r["quantity"], r["comp"], r["fot"], r["overhead"], r["cost"]]
        add_row(ws, [*values, round(r["profit_percent"], 2), r["profit"], r["final"]], money, len(headers))
    t = summary["totals"]
    values = ["Итого", None, t["comp"], t["fot"], t["overhead"], t["cost"], round(t["profit_percent"], 2)]
    add_row(ws, [*values, t["profit"], t["final"]], money, len(headers), is_total=True)
    add_row(ws, [f"Итого с НДС {summary['nds_percent']:g}%", *[None] * 7, t["final_nds"]], (9,), len(headers), True)
    auto_fit_columns(ws)

    # Модули: строка раздела (кол-во и стоимость без НДС — у всех разделов, включая корпуса и работы),
    # под ней модули раздела
    ws = wb.create_sheet("Комплектация")
    headers = ["Раздел", "ID", "Модуль", "Кол-во", "Цена за ед.", "Сумма"]
    sheet_title(ws, f"{title} — комплектация", len(headers))
    write_header(ws, headers)
    modules_by_section: dict[str, list[dict]] = {}
    for r in kit["modules"]:
        modules_by_section.setdefault(r["section"], []).append(r)
    for s in summary["rows"]:
        unit = s["final"] / s["quantity"] if s["quantity"] else s["final"]
        add_row(ws, [s["name"], None, None, s["quantity"], unit, s["final"]], (5, 6), 6, is_total=True)
        for r in modules_by_section.pop(s["name"], []):
            add_row(ws, [r["section"], r["module_id"], r["name"], r["quantity"], r["price"], r["sum"]], (5, 6), 6)
    for rows in modules_by_section.values():
        for r in rows:
            add_row(ws, [r["section"], r["module_id"], r["name"], r["quantity"], r["price"], r["sum"]], (5, 6), 6)
    add_row(ws, ["Итого модули", None, None, None, None, sum(r["sum"] for r in kit["modules"])], (6,), 6, True)
    add_row(ws, ["Итого без НДС", None, None, None, None, summary["totals"]["final"]], (6,), 6, True)
    auto_fit_columns(ws)

    # Материалы (подмодули раскрыты)
    ws = wb.create_sheet("Материалы")
    headers = ["ID", "Код 1С", "Тип материала", "Название", "Ном Ток", "Цена за ед.", "Кол-во", "Сумма"]
    sheet_title(ws, f"{title} — сводная по материалам", len(headers))
    write_header(ws, headers)
    for r in kit["materials"]:
        values = [r["id"], r["code_1c"], r["type"], r["name"], r["nom_tok"] or None, r["price"], r["quantity"]]
        add_row(ws, [*values, r["sum"]], (6, 8), len(headers))
    add_row(ws, ["Итого", *[None] * 6, sum(r["sum"] for r in kit["materials"])], (8,), len(headers), True)
    auto_fit_columns(ws)

    return xlsx_response(wb, f"komplektaciya_tkp_{req.id}_{dt.now().strftime('%Y%m%d')}.xlsx")


@mt_calc_router.get("/api/requests/{req_id}/tkp-kit/export-excel")
async def export_tkp_kit_excel(req_id: int, session: AsyncSession = Depends(db_helper.session_dependency)):
    """Комплектация ТКП таблицей, как в файле ТКП: №, наименование (блок и состав), кол-во, цена, стоимость,
    итого, НДС, всего — с формулами."""
    from datetime import datetime as dt

    from openpyxl.styles import Alignment, Font

    from scr.dbase import crud_requests, crud_settings
    from scr.dbase.crud_tkp import build_tkp_kit
    from scr.utils.excel import MONEY_FMT, border_row, write_header, xlsx_response

    req = await crud_requests.get_request_by_id(session, req_id)
    if not req:
        raise HTTPException(status_code=404, detail="ТКП не найдена")
    if await crud_module_slots.sync_hidden_slots(session, req_id):
        await session.commit()
    kit = await build_tkp_kit(session, req)
    nds_val = await crud_settings.get_setting(session, "nds_percent")
    nds = float(nds_val) if nds_val else 20.0

    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    wb = Workbook()
    ws = wb.active
    ws.title = "Комплектация"
    ws.append([f"Комплектация ТКП № {req.tkp_num or req.id}"])
    ws.cell(1, 1).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
    write_header(ws, ["№", "Наименование оборудования", "Кол-во, шт", "Цена за ед продукции без НДС, руб",
                      "Стоимость без НДС, руб"])  # fmt: skip
    for col in range(1, 6):
        ws.cell(ws.max_row, col).alignment = Alignment(wrap_text=True, vertical="center")
    first = ws.max_row + 1
    n = 0
    for section in kit:
        if section["header"]:
            ws.append([None, f"{section['name']} в составе:"])
            ws.cell(ws.max_row, 2).font = bold
            border_row(ws, ws.max_row, 5)
        for block in section["blocks"]:
            n += 1
            # Состав блока — в той же ячейке, построчно
            text = ", \n".join([block["title"], *block["lines"]])
            ws.append([n, text, block["quantity"], block["price"]])
            r = ws.max_row
            ws.cell(r, 5).value = f"=D{r}*C{r}"
            ws.cell(r, 2).font = bold
            for col in range(1, 6):
                ws.cell(r, col).alignment = wrap
            ws.cell(r, 4).number_format = MONEY_FMT
            ws.cell(r, 5).number_format = MONEY_FMT
            border_row(ws, r, 5)
    last = ws.max_row
    for label, formula in (
        ("Итого", f"=SUM(E{first}:E{last})" if last >= first else 0),
        (f"НДС {nds:g}%:", f"=E{last + 1}*{nds / 100:g}"),
        ("Всего:", f"=E{last + 1}+E{last + 2}"),
    ):
        ws.append([None, None, None, label, formula])
        ws.cell(ws.max_row, 4).font = bold
        ws.cell(ws.max_row, 5).number_format = MONEY_FMT
    for col, width in zip("ABCDE", (5, 70, 10, 18, 18), strict=True):
        ws.column_dimensions[col].width = width
    return xlsx_response(wb, f"komplektaciya_tkp_{req.tkp_num or req.id}_{dt.now().strftime('%Y%m%d')}.xlsx")


@mt_calc_router.get("/api/requests/{req_id}/calc/{section}/export-excel")
async def export_calc_excel(req_id: int, section: str, session: AsyncSession = Depends(db_helper.session_dependency)):
    """Сводная по материалам расчёта раздела: по столбцам — блоки (кол-во и сумма), справа — итого,
    внизу — итоги по столбцам, ФОТ, накладные, рентабельность и конечная цена."""
    from datetime import datetime as dt

    from openpyxl.styles import Alignment, Font

    from scr.dbase import crud_requests
    from scr.dbase.crud_tkp import build_calc_materials
    from scr.utils.excel import MONEY_FMT, auto_fit_columns, border_row, safe_filename, xlsx_response

    req = await crud_requests.get_request_by_id(session, req_id)
    if not req:
        raise HTTPException(status_code=404, detail="ТКП не найдена")
    mt = await crud_module_types.get_module_type_by_slug(session, section)
    if not mt:
        raise HTTPException(status_code=404, detail="Раздел не найден")
    if await crud_module_slots.sync_hidden_slots(session, req_id):
        await session.commit()
    calcs = await crud_module_types.get_calcs(session, req_id, mt.id)
    materials, per_calc = await build_calc_materials(session, calcs)
    multi = section in crud_module_types.MULTI_BLOCK_SECTIONS
    split = section in crud_module_types.SPLIT_QTY_SECTIONS
    works = section in crud_module_types.WORKS_SECTIONS
    block_qty = [(c.quantity if c.quantity is not None else 1) if multi else 1 for c in calcs]

    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    fixed = 4  # ID, Материал, Цена за ед., Дата цены
    total_col = fixed + 2 * len(calcs) + 1
    cols = total_col + 1

    wb = Workbook()
    ws = wb.active
    ws.title = safe_filename(mt.name)[:31] or "Расчёт"
    org = req.company.name if req.company else ""
    title = f"Расчёт: {mt.name} · ТКП № {req.tkp_num or req.id}" + (f" · {org}" if org else "")
    ws.append([title])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    ws.cell(1, 1).font = Font(bold=True, size=14)

    # Двухуровневая шапка: блок → Кол-во / Сумма
    ws.append([])
    ws.append([])
    for col, name in enumerate(["ID", "Материал", "Цена за ед.", "Дата цены"], start=1):
        ws.cell(2, col, name)
        ws.merge_cells(start_row=2, start_column=col, end_row=3, end_column=col)
    for i, calc in enumerate(calcs):
        label = f"{mt.name} {i + 1}" if multi else mt.name
        if split:
            label += f" (1СШ: {calc.qty_1sh or 0}, 2СШ: {calc.qty_2sh or 0})"
        elif multi:
            label += f" × {block_qty[i]}"
        col = fixed + 1 + 2 * i
        ws.cell(2, col, label)
        ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + 1)
        ws.cell(3, col, "Кол-во")
        ws.cell(3, col + 1, "Сумма")
    ws.cell(2, total_col, "Итого" + (" (с учётом кол-ва блоков)" if multi else ""))
    ws.merge_cells(start_row=2, start_column=total_col, end_row=2, end_column=total_col + 1)
    ws.cell(3, total_col, "Кол-во")
    ws.cell(3, total_col + 1, "Сумма")
    from scr.utils.excel import HEADER_FILL

    for row in (2, 3):
        for col in range(1, cols + 1):
            c = ws.cell(row, col)
            c.font = bold
            c.fill = HEADER_FILL
            c.alignment = center
        border_row(ws, row, cols)

    money_cols = [3] + [fixed + 2 + 2 * i for i in range(len(calcs))] + [total_col + 1]

    def add_row(values: dict[int, object], is_total: bool = False) -> int:
        ws.append([values.get(col) for col in range(1, cols + 1)])
        row_num = ws.max_row
        for col in money_cols:
            ws.cell(row_num, col).number_format = MONEY_FMT
        border_row(ws, row_num, cols)
        if is_total:
            for col in range(1, cols + 1):
                ws.cell(row_num, col).font = bold
        return row_num

    # Материалы
    block_sums = [0.0] * len(calcs)
    for mat in materials:
        values: dict[int, object] = {1: mat["id"], 2: mat["name"], 3: mat["price"], 4: mat["date"]}
        total_qty = 0.0
        for i, acc in enumerate(per_calc):
            qty = acc.get(mat["id"], 0)
            if qty:
                values[fixed + 1 + 2 * i] = qty
                values[fixed + 2 + 2 * i] = qty * mat["price"]
                block_sums[i] += qty * mat["price"]
            total_qty += qty * block_qty[i]
        values[total_col] = total_qty
        values[total_col + 1] = total_qty * mat["price"]
        row_num = add_row(values)
        ws.cell(row_num, 4).number_format = "DD.MM.YYYY"

    def footer(label: str, per_block: list[float], qty_values: list | None = None, qty_total=None) -> None:
        """Строка итогов: значение по блоку — в колонку «Сумма», итого — с учётом кол-ва блоков."""
        values: dict[int, object] = {2: label}
        for i, value in enumerate(per_block):
            values[fixed + 2 + 2 * i] = value
            if qty_values is not None:
                values[fixed + 1 + 2 * i] = qty_values[i]
        values[total_col + 1] = sum(v * q for v, q in zip(per_block, block_qty, strict=True))
        if qty_total is not None:
            values[total_col] = qty_total
        add_row(values, is_total=True)

    fots = [float(c.fot or 0) for c in calcs]
    overheads = [float(c.overhead or 0) for c in calcs]
    costs = [float(c.total_cost or 0) for c in calcs]
    finals = [float(c.final_price or 0) for c in calcs]
    profits = [f - c for f, c in zip(finals, costs, strict=True)]
    pcts = [round(p / c * 100, 2) if c else 0.0 for p, c in zip(profits, costs, strict=True)]
    total_cost = sum(c * q for c, q in zip(costs, block_qty, strict=True))
    total_profit = sum(p * q for p, q in zip(profits, block_qty, strict=True))

    footer("Итого материалы", block_sums)
    if works:
        footer("Себестоимость работ", [float(c.works_cost or 0) for c in calcs])
    footer("ФОТ", fots)
    footer("Накладные", overheads)
    footer("Себестоимость", costs)
    total_pct = round(total_profit / total_cost * 100, 2) if total_cost else 0.0
    footer("Рентабельность (% / сумма)", profits, pcts, total_pct)
    footer("Конечная цена" + (" за блок" if multi else ""), finals)
    if multi:
        values = {2: "Количество блоков", total_col: sum(block_qty)}
        for i, q in enumerate(block_qty):
            values[fixed + 1 + 2 * i] = q
        add_row(values, is_total=True)
        values = {2: "Конечная цена × кол-во"}
        for i, (f, q) in enumerate(zip(finals, block_qty, strict=True)):
            values[fixed + 2 + 2 * i] = f * q
        values[total_col + 1] = sum(f * q for f, q in zip(finals, block_qty, strict=True))
        add_row(values, is_total=True)

    auto_fit_columns(ws)
    ws.column_dimensions["A"].width = 8
    for i in range(len(calcs) + 1):
        ws.column_dimensions[ws.cell(3, fixed + 1 + 2 * i).column_letter].width = 10
        ws.column_dimensions[ws.cell(3, fixed + 2 + 2 * i).column_letter].width = 16
    ws.row_dimensions[2].height = 32
    ws.freeze_panes = ws.cell(4, fixed + 1)
    name = f"raschet_{safe_filename(mt.slug)}_tkp_{req.id}_{dt.now().strftime('%Y%m%d')}.xlsx"
    return xlsx_response(wb, name)
