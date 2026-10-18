from fastapi import APIRouter, Depends
from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession

from scr.dbase import crud_payments
from scr.dbase.database import db_helper
from scr.dbase.schemas.schemas import (
    PaymentItemResponseSchema,
    PaymentItemUpdateSchema,
)

pay_router = APIRouter(prefix="/api/payments", tags=["Payments"])


@pay_router.get("/by-request/{request_id}", response_model=list[PaymentItemResponseSchema])
async def read_payments(request_id: int, session: AsyncSession = Depends(db_helper.session_dependency)):
    items = await crud_payments.ensure_payment_items(session, request_id)
    await session.commit()
    return items


@pay_router.put("/{item_id}", response_model=PaymentItemResponseSchema)
async def update_payment(
    item_id: int,
    data: PaymentItemUpdateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    data.id = item_id
    result = await crud_payments.update_payment_item(session, data)
    if not result:
        from fastapi import HTTPException

        raise HTTPException(status_code=404, detail="Платёж не найден")
    await session.commit()
    return result


class PaymentRegistrySchema(BaseModel):
    """ТКП, видимые в реестре с учётом фильтров, и описание фильтра для заголовка"""

    ids: list[int]
    filter_text: str = ""


@pay_router.post("/registry-excel")
async def payment_registry_excel(
    data: PaymentRegistrySchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    """Реестр оплат: порядок оплаты по ТКП с разбивкой по номерам ТКП и сводная по видам платежей."""
    from datetime import datetime as dt

    from openpyxl import Workbook
    from openpyxl.styles import Font

    from scr.constants import PAYMENT_TYPES
    from scr.utils.excel import MONEY_FMT, auto_fit_columns, border_row, write_header, xlsx_response

    requests = await crud_payments.get_requests_with_payments(session, data.ids)
    today = dt.now()
    bold = Font(bold=True)
    title = "Реестр оплат" + (f" — фильтр: {data.filter_text}" if data.filter_text else "")

    def money(v) -> float:
        return round(float(v or 0), 2)

    def add_row(ws, values: list, money_cols: tuple, is_total: bool = False) -> int:
        ws.append(values)
        row_num = ws.max_row
        for col in money_cols:
            ws.cell(row_num, col).number_format = MONEY_FMT
        border_row(ws, row_num, len(values))
        if is_total:
            for col in range(1, len(values) + 1):
                ws.cell(row_num, col).font = bold
        return row_num

    def start_sheet(ws, headers: list[str]) -> None:
        ws.append([title])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws.cell(1, 1).font = Font(bold=True, size=14)
        write_header(ws, headers)

    wb = Workbook()

    # Подробно: строки порядка оплаты, сгруппированные по ТКП
    ws = wb.active
    ws.title = "Реестр оплат"
    headers = [
        "№ ТКП", "Номер заказа", "Организация", "Стоимость с НДС", "Вид платежа", "%", "Сумма",
        "Оплатить до", "Оплачено", "Дата оплаты", "Остаток",
    ]  # fmt: skip
    start_sheet(ws, headers)
    money_cols = (4, 7, 9, 11)
    by_type = {pt: [0.0, 0.0] for pt in PAYMENT_TYPES}
    per_request = []
    grand = [0.0, 0.0]
    for req in requests:
        items = [
            pi
            for pi in crud_payments.sorted_payment_items(req.payment_items)
            if money(pi.amount) or money(pi.paid_amount)
        ]
        if not items:
            continue
        org = req.company.name if req.company else ""
        req_total = [0.0, 0.0]
        overdue = 0.0
        for pi in items:
            amount, paid = money(pi.amount), money(pi.paid_amount)
            rest = round(amount - paid, 2)
            row_num = add_row(
                ws,
                [
                    req.tkp_num or str(req.id), req.factory_order_num or "", org, money(req.cost), pi.payment_type,
                    money(pi.percent), amount, pi.due_date, paid, pi.paid_date, rest,
                ],
                money_cols,
            )  # fmt: skip
            for col in (8, 10):
                ws.cell(row_num, col).number_format = "DD.MM.YYYY"
            req_total[0] += amount
            req_total[1] += paid
            if pi.due_date and pi.due_date < today and rest > 0:
                overdue += rest
            totals = by_type.setdefault(pi.payment_type, [0.0, 0.0])
            totals[0] += amount
            totals[1] += paid
        rest_total = round(req_total[0] - req_total[1], 2)
        add_row(
            ws,
            [f"Итого по ТКП {req.tkp_num or req.id}", req.factory_order_num or "", None, None, None, None,
             req_total[0], None, req_total[1], None, rest_total],
            money_cols,
            is_total=True,
        )  # fmt: skip
        per_request.append((req, org, req_total, rest_total, overdue))
        grand[0] += req_total[0]
        grand[1] += req_total[1]
    add_row(
        ws,
        ["Итого", None, None, None, None, None, grand[0], None, grand[1], None, round(grand[0] - grand[1], 2)],
        money_cols,
        is_total=True,
    )
    auto_fit_columns(ws)

    # По ТКП: одна строка на ТКП
    ws = wb.create_sheet("По ТКП")
    headers = ["№ ТКП", "Номер заказа", "Организация", "Стоимость с НДС", "К оплате", "Оплачено", "Остаток",
               "Просрочено"]  # fmt: skip
    start_sheet(ws, headers)
    for req, org, req_total, rest_total, overdue in per_request:
        add_row(
            ws,
            [req.tkp_num or str(req.id), req.factory_order_num or "", org, money(req.cost), req_total[0],
             req_total[1], rest_total, round(overdue, 2)],
            (4, 5, 6, 7, 8),
        )  # fmt: skip
    add_row(
        ws,
        ["Итого", None, None, sum(money(r[0].cost) for r in per_request), grand[0], grand[1],
         round(grand[0] - grand[1], 2), round(sum(r[4] for r in per_request), 2)],
        (4, 5, 6, 7, 8),
        is_total=True,
    )  # fmt: skip
    auto_fit_columns(ws)

    # Сводная по видам платежей
    ws = wb.create_sheet("По видам платежей")
    headers = ["Вид платежа", "К оплате", "Оплачено", "Остаток"]
    start_sheet(ws, headers)
    for pt, (amount, paid) in by_type.items():
        if amount or paid:
            add_row(ws, [pt, amount, paid, round(amount - paid, 2)], (2, 3, 4))
    add_row(ws, ["Итого", grand[0], grand[1], round(grand[0] - grand[1], 2)], (2, 3, 4), is_total=True)
    auto_fit_columns(ws)

    return xlsx_response(wb, f"reestr_oplat_{today.strftime('%Y%m%d')}.xlsx")
