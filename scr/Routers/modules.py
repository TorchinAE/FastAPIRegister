import io

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from scr.dbase import crud_modules
from scr.dbase.database import db_helper
from scr.dbase.schemas.schemas import (
    MaterialResponseSchema,
    ModuleCreateSchema,
    ModuleFullResponseSchema,
    ModuleItemCreateSchema,
    ModuleItemResponseSchema,
    ModuleResponseSchema,
    ModuleUpdateSchema,
    PaginatedResponse,
)

mod_router = APIRouter(prefix="/api/modules", tags=["Modules"])


def _module_to_response(mod) -> ModuleResponseSchema:
    return ModuleResponseSchema(
        id=mod.id,
        name=mod.name,
        items_count=mod.items_count,
        total_price=mod.total_price,
        nom_tok=mod.nom_tok or 0,
        nom_tok_manual=bool(mod.nom_tok_manual),
        created_by=mod.created_by,
    )


def _module_item_to_response(item) -> ModuleItemResponseSchema:
    return ModuleItemResponseSchema(
        id=item.id,
        module_id=item.module_id,
        material_id=item.material_id,
        sub_module_id=item.sub_module_id,
        quantity=item.quantity,
        material=MaterialResponseSchema.model_validate(item.material) if item.material else None,
        sub_module_name=item.sub_module.name if item.sub_module else None,
    )


@mod_router.post("/", response_model=ModuleResponseSchema)
async def add_module(
    data: ModuleCreateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    new_mod = await crud_modules.add_module(session=session, in_mod=data)
    await session.commit()
    await session.refresh(new_mod)
    return _module_to_response(new_mod)


@mod_router.get("/", response_model=PaginatedResponse)
async def read_modules(
    search: str | None = Query(None),
    page: int = Query(1, ge=1),
    per_page: int = Query(20, ge=1, le=100),
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    items, total = await crud_modules.get_modules(session=session, search=search, page=page, per_page=per_page)
    return PaginatedResponse(
        items=[_module_to_response(i) for i in items],
        total=total,
        page=page,
        per_page=per_page,
        pages=(total + per_page - 1) // per_page,
    )


@mod_router.get("/all", response_model=list[ModuleResponseSchema])
async def read_all_modules(
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    items = await crud_modules.get_all_modules(session=session)
    return [_module_to_response(i) for i in items]


def _auto_fit_columns(ws):
    from openpyxl.utils import get_column_letter

    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value is not None:
                cell_len = len(str(cell.value))
                max_len = max(max_len, cell_len)
        ws.column_dimensions[col_letter].width = min(max_len + 3, 60)


@mod_router.get("/export-excel")
async def export_modules_excel(
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    from datetime import datetime as dt

    from openpyxl.styles import Border, Font, PatternFill, Side

    modules = await crud_modules.get_all_modules(session=session)
    wb = Workbook()

    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    def _style_header(ws, cols, row_num=1):
        for col_idx in range(1, cols + 1):
            cell = ws.cell(row_num, col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border

    def _style_row(ws, row_num, cols):
        for col_idx in range(1, cols + 1):
            ws.cell(row_num, col_idx).border = thin_border

    ws1 = wb.active
    ws1.title = "Модули"
    ws1.append(["ID", "Название", "Компонентов", "Стоимость"])
    _style_header(ws1, 4)
    for mod in modules:
        ws1.append([mod.id, mod.name, mod.items_count, float(mod.total_price)])
        _style_row(ws1, ws1.max_row, 4)
        ws1.cell(ws1.max_row, 4).number_format = "# ##0.00"
    _auto_fit_columns(ws1)

    ws2 = wb.create_sheet("Состав")
    ws2.append(["Модуль", "Тип компонента", "Название компонента", "Количество"])
    _style_header(ws2, 4)
    for mod in modules:
        items = await crud_modules.get_module_items(session, mod.id)
        for item in items:
            if item.material:
                ws2.append([mod.name, "материал", item.material.name, item.quantity])
            elif item.sub_module:
                ws2.append([mod.name, "модуль", item.sub_module.name, item.quantity])
            _style_row(ws2, ws2.max_row, 4)
    _auto_fit_columns(ws2)

    from urllib.parse import quote

    date_str = dt.now().strftime("%Y%m%d")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"modules_{date_str}.xlsx"
    encoded = quote(filename)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded}"},
    )


@mod_router.post("/import-excel")
async def import_modules_excel(
    file: UploadFile = File(...),
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    content = await file.read()
    wb = load_workbook(io.BytesIO(content))

    imported = 0
    if "Модули" in wb.sheetnames:
        ws = wb["Модули"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            name = str(row[0]).strip() if row[0] else None
            if not name:
                continue
            await crud_modules.add_module(session, ModuleCreateSchema(name=name))
            imported += 1

    if "Состав" in wb.sheetnames:
        ws = wb["Состав"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            mod_name = str(row[0]).strip()
            mod = await crud_modules.get_module_by_name(session, mod_name)
            if not mod:
                continue
            item_type = str(row[1]).strip() if row[1] else ""
            item_name = str(row[2]).strip() if row[2] else ""
            quantity = int(row[3]) if row[3] else 1
            if item_type == "материал":
                from scr.dbase import crud_materials

                mat = await crud_materials.get_material_by_name(session, item_name)
                if mat:
                    await crud_modules.add_module_item(
                        session,
                        mod.id,
                        ModuleItemCreateSchema(material_id=mat.id, quantity=quantity),
                    )
            elif item_type == "модуль":
                sub_mod = await crud_modules.get_module_by_name(session, item_name)
                if sub_mod:
                    await crud_modules.add_module_item(
                        session,
                        mod.id,
                        ModuleItemCreateSchema(sub_module_id=sub_mod.id, quantity=quantity),
                    )

    await crud_modules.recalc_nom_tok(session)
    await session.commit()
    return {"message": f"Импортировано модулей: {imported}", "imported": imported}


MODULE_ITEMS_HEADERS = ["Тип", "ID", "Код 1С", "Тип материала", "Название", "Цена за ед.", "Кол-во", "Сумма"]

_MODULE_ITEMS_ALIASES = {
    "тип": "kind",
    "id": "id",
    "id материала": "id",
    "код 1с": "code_1c",
    "код 1c": "code_1c",
    "тип материала": "mat_type",
    "название": "name",
    "наименование": "name",
    "материал": "name",
    "цена за ед.": "price",
    "цена за ед": "price",
    "цена": "price",
    "кол-во": "qty",
    "количество": "qty",
    "кол.": "qty",
}


@mod_router.get("/{mod_id}/export-excel")
async def export_module_excel(mod_id: int, session: AsyncSession = Depends(db_helper.session_dependency)):
    from datetime import datetime as dt

    from openpyxl.styles import Font

    from scr.utils.excel import MONEY_FMT, auto_fit_columns, border_row, safe_filename, write_header, xlsx_response

    mod = await crud_modules.get_module_by_id(session=session, mod_id=mod_id)
    if not mod:
        raise HTTPException(status_code=404, detail="Модуль не найден")
    prices = await crud_modules.get_module_prices(session)
    all_modules = {m.id: m for m in await crud_modules.get_all_modules(session)}

    def _bom_flat(module, qty=1, stack=frozenset()):
        result = []
        for it in module.items:
            if it.material:
                result.append(
                    {"name": it.material.name, "price": float(it.material.price), "quantity": it.quantity * qty}
                )
            elif it.sub_module_id in all_modules and it.sub_module_id not in stack:
                sub = all_modules[it.sub_module_id]
                result.extend(_bom_flat(sub, it.quantity * qty, stack | {module.id}))
        return result

    wb = Workbook()
    ws = wb.active
    ws.title = "Состав модуля"
    cols = len(MODULE_ITEMS_HEADERS)

    # Row 1: module name
    ws.append([mod.name])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    ws.cell(1, 1).font = Font(bold=True, size=14)

    # Row 2: headers
    write_header(ws, MODULE_ITEMS_HEADERS)

    # Rows 3+: items (F — цена, G — кол-во, H — сумма)
    for item in mod.items:
        if item.material:
            mat = item.material
            await session.refresh(mat, ["type"])
            row_data = [
                "Материал",
                mat.id,
                mat.code_1c or "",
                mat.type.name if mat.type else "",
                mat.name,
                float(mat.price),
                item.quantity,
            ]
        elif item.sub_module:
            row_data = [
                "Модуль",
                item.sub_module.id,
                "",
                "",
                item.sub_module.name,
                prices.get(item.sub_module.id, 0.0),
                item.quantity,
            ]
        else:
            continue
        ws.append(row_data)
        row_num = ws.max_row
        ws.cell(row_num, 8).value = f"=F{row_num}*G{row_num}"
        ws.cell(row_num, 6).number_format = MONEY_FMT
        ws.cell(row_num, 8).number_format = MONEY_FMT
        border_row(ws, row_num, cols)

    # Total row: отделена пустой строкой, чтобы импорт остановился на составе
    items_end = ws.max_row
    ws.append([])
    ws.append([""] * (cols - 2) + ["Итого:", f"=SUM(H3:H{items_end})"])
    total_row = ws.max_row
    ws.cell(total_row, cols - 1).font = Font(bold=True)
    ws.cell(total_row, cols).font = Font(bold=True)
    ws.cell(total_row, cols).number_format = MONEY_FMT

    # BOM section
    bom_items = _bom_flat(mod)
    if bom_items:
        ws.append([])
        ws.append(["Раскрытие состава (BOM)"])
        bom_title_row = ws.max_row
        ws.merge_cells(start_row=bom_title_row, start_column=1, end_row=bom_title_row, end_column=5)
        ws.cell(bom_title_row, 1).font = Font(bold=True, size=12)

        write_header(ws, ["", "Материал", "Цена за ед.", "Кол-во", "Сумма"])
        bom_hdr_row = ws.max_row
        for bi in bom_items:
            ws.append(["", bi["name"], bi["price"], bi["quantity"], None])
            row_num = ws.max_row
            ws.cell(row_num, 3).number_format = MONEY_FMT
            ws.cell(row_num, 5).number_format = MONEY_FMT
            ws.cell(row_num, 5).value = f"=C{row_num}*D{row_num}"
            border_row(ws, row_num, 5)

        ws.append(["", "", "", "Итого:", f"=SUM(E{bom_hdr_row + 1}:E{ws.max_row})"])
        bom_total_row = ws.max_row
        ws.cell(bom_total_row, 4).font = Font(bold=True)
        ws.cell(bom_total_row, 5).font = Font(bold=True)
        ws.cell(bom_total_row, 5).number_format = MONEY_FMT
        border_row(ws, bom_total_row, 5)

    auto_fit_columns(ws)
    return xlsx_response(wb, f"{safe_filename(mod.name)}_{dt.now().strftime('%Y%m%d')}.xlsx")


async def _find_material(session: AsyncSession, mat_id: int | None, code_1c: str | None, name: str | None):
    """Поиск материала: ID → Код 1С → Название."""
    from scr.dbase import crud_materials

    if mat_id:
        mat = await crud_materials.get_material_by_id(session, mat_id)
        if mat:
            return mat
    if code_1c:
        mat = await crud_materials.get_material_by_code_1c(session, code_1c)
        if mat:
            return mat
    if name:
        return await crud_materials.get_material_by_name(session, name)
    return None


@mod_router.post("/{mod_id}/import-excel")
async def import_module_items_excel(
    mod_id: int,
    file: UploadFile = File(...),
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    """Загрузка состава модуля. Колонки: Название, Цена за ед., Кол-во, Тип материала (+ необязательные ID, Код 1С, Тип).

    Материал ищется по ID, затем по Коду 1С, затем по названию; не найден — создаётся новый
    (тип по умолчанию «Производство»). Если компонент уже есть в модуле — количество заменяется.
    """
    from datetime import date as date_cls

    from scr.dbase import crud_material_types, crud_materials
    from scr.dbase.models import ModuleItem
    from scr.dbase.schemas.schemas import MaterialCreateSchema, MaterialTypeCreateSchema
    from scr.utils.excel import cell, data_rows, find_header, normalize, to_float, to_int

    mod = await crud_modules.get_module_by_id(session, mod_id)
    if not mod:
        raise HTTPException(status_code=404, detail="Модуль не найден")
    try:
        wb = load_workbook(io.BytesIO(await file.read()), data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Не удалось прочитать файл Excel") from exc
    rows = list(wb.active.iter_rows(values_only=True))
    header_idx, col_map = find_header(rows, _MODULE_ITEMS_ALIASES, required="name")
    if header_idx < 0:
        raise HTTPException(status_code=400, detail="Не найдена строка заголовка с колонкой «Название»")

    types, _ = await crud_material_types.get_all_types(session, per_page=10000)
    types_map = {t.name.strip().lower(): t.id for t in types}

    async def type_id_for(type_name) -> int:
        name = str(type_name).strip() if type_name else "Производство"
        key = name.lower()
        if key not in types_map:
            new_type = await crud_material_types.add_type(session, MaterialTypeCreateSchema(name=name))
            types_map[key] = new_type.id
        return types_map[key]

    existing = {(it.material_id, it.sub_module_id): it for it in mod.items}
    added = updated = created_materials = 0
    skipped: list[str] = []

    for row in data_rows(rows, header_idx):
        name = cell(row, col_map, "name")
        name = str(name).strip() if name is not None else None
        row_id = to_int(cell(row, col_map, "id"))
        if not name and not row_id:
            continue
        qty = to_int(cell(row, col_map, "qty"), 1)
        if qty is None or qty < 1:
            skipped.append(f"{name}: количество {qty}")
            continue

        if normalize(cell(row, col_map, "kind")) == "модуль":
            sub = await crud_modules.get_module_by_id(session, row_id) if row_id else None
            if not sub and name:
                sub = await crud_modules.get_module_by_name(session, name)
            if not sub or sub.id == mod_id:
                skipped.append(f"{name}: модуль не найден")
                continue
            key = (None, sub.id)
        else:
            code_1c = cell(row, col_map, "code_1c")
            code_1c = str(code_1c).strip() if code_1c is not None else None
            mat = await _find_material(session, row_id, code_1c, name)
            if not mat:
                if not name:
                    skipped.append(f"ID {row_id}: материал не найден")
                    continue
                mat = await crud_materials.add_material(
                    session,
                    MaterialCreateSchema(
                        name=name,
                        price=max(to_float(cell(row, col_map, "price")), 0),
                        code_1c=code_1c,
                        type_id=await type_id_for(cell(row, col_map, "mat_type")),
                        stats=False,
                        ruchn=False,
                        date=date_cls.today(),
                    ),
                )
                created_materials += 1
            key = (mat.id, None)

        if key in existing:
            existing[key].quantity = qty
            updated += 1
        else:
            item = ModuleItem(module_id=mod_id, material_id=key[0], sub_module_id=key[1], quantity=qty)
            session.add(item)
            existing[key] = item
            added += 1

    await session.flush()
    await crud_modules.recalc_nom_tok(session)
    await session.commit()
    parts = [f"Добавлено: {added}", f"обновлено: {updated}", f"новых материалов: {created_materials}"]
    if skipped:
        parts.append(f"пропущено: {len(skipped)}")
    return {
        "message": ", ".join(parts),
        "added": added,
        "updated": updated,
        "created_materials": created_materials,
        "skipped": skipped,
    }


@mod_router.post("/{mod_id}/duplicate", response_model=ModuleFullResponseSchema)
async def duplicate_module(mod_id: int, session: AsyncSession = Depends(db_helper.session_dependency)):
    mod = await crud_modules.get_module_by_id(session=session, mod_id=mod_id)
    if not mod:
        raise HTTPException(status_code=404, detail="Модуль не найден")
    new_name = f"{mod.name} - копия"
    new_mod = await crud_modules.add_module(session=session, in_mod=ModuleCreateSchema(name=new_name))
    await session.flush()
    for item in mod.items:
        await crud_modules.add_module_item(
            session,
            new_mod.id,
            ModuleItemCreateSchema(
                material_id=item.material_id,
                sub_module_id=item.sub_module_id,
                quantity=item.quantity,
            ),
        )
    await session.commit()
    new_mod = await crud_modules.get_module_by_id(session=session, mod_id=new_mod.id)
    items_resp = [_module_item_to_response(i) for i in new_mod.items]
    return ModuleFullResponseSchema(
        id=new_mod.id,
        name=new_mod.name,
        items=items_resp,
        total_price=new_mod.total_price,
        created_by=new_mod.created_by,
    )


@mod_router.get("/{mod_id}", response_model=ModuleFullResponseSchema)
async def read_module(mod_id: int, session: AsyncSession = Depends(db_helper.session_dependency)):
    mod = await crud_modules.get_module_by_id(session=session, mod_id=mod_id)
    if not mod:
        raise HTTPException(status_code=404, detail="Модуль не найден")
    items_resp = [_module_item_to_response(i) for i in mod.items]
    return ModuleFullResponseSchema(
        id=mod.id,
        name=mod.name,
        items=items_resp,
        total_price=mod.total_price,
        created_by=mod.created_by,
    )


@mod_router.patch("/", response_model=ModuleResponseSchema)
async def update_module(
    data: ModuleUpdateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    result = await crud_modules.update_module(session=session, upd_mod=data)
    if not result:
        raise HTTPException(status_code=404, detail="Модуль не найден")
    if data.nom_tok:
        from scr.dbase import crud_nominal_currents

        await crud_nominal_currents.ensure_value(session, data.nom_tok)
    await session.commit()
    await session.refresh(result)
    return _module_to_response(result)


@mod_router.delete("/{mod_id}")
async def delete_module(mod_id: int, session: AsyncSession = Depends(db_helper.session_dependency)):
    result = await crud_modules.delete_module(session=session, mod_id=mod_id)
    if not result:
        raise HTTPException(status_code=404, detail="Модуль не найден")
    await session.commit()
    return {"message": "Модуль удалён"}


@mod_router.post("/{mod_id}/items", response_model=ModuleItemResponseSchema)
async def add_module_item(
    mod_id: int,
    data: ModuleItemCreateSchema,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    item = await crud_modules.add_module_item(session, mod_id, data)
    if not item:
        raise HTTPException(status_code=404, detail="Модуль не найден")
    await crud_modules.recalc_nom_tok(session)
    await session.commit()
    await session.refresh(item)
    # Eagerly load relationships after commit (they get expired)
    if item.material_id:
        await session.refresh(item, ["material"])
    if item.sub_module_id:
        await session.refresh(item, ["sub_module"])
    return _module_item_to_response(item)


@mod_router.patch("/items/{item_id}", response_model=ModuleItemResponseSchema)
async def update_module_item(
    item_id: int,
    data: dict,
    session: AsyncSession = Depends(db_helper.session_dependency),
):
    quantity = data.get("quantity", 1)
    item = await crud_modules.update_module_item_quantity(session, item_id, quantity)
    if not item:
        raise HTTPException(status_code=404, detail="Элемент не найден")
    await session.commit()
    await session.refresh(item)
    if item.material_id:
        await session.refresh(item, ["material"])
    if item.sub_module_id:
        await session.refresh(item, ["sub_module"])
    return _module_item_to_response(item)


@mod_router.delete("/items/{item_id}")
async def delete_module_item(item_id: int, session: AsyncSession = Depends(db_helper.session_dependency)):
    result = await crud_modules.delete_module_item(session, item_id)
    if not result:
        raise HTTPException(status_code=404, detail="Элемент не найден")
    await crud_modules.recalc_nom_tok(session)
    await session.commit()
    return {"message": "Элемент удалён"}
