"""ТКП в Excel по шаблону xlsx_templates/ТКП.xlsx (шапка с логотипом и реквизитами, таблица комплектации, условия).

Шаблон: строка 14 — образец строки блока, 15 — образец заголовка «… в составе:», 16–18 — Итого / НДС / Всего.
Строки комплектации вставляются вместо двух образцов, всё ниже сдвигается.
openpyxl при сохранении теряет картинки, поэтому логотип (drawing) переносится из шаблона в готовый файл.
"""

import io
import math
import re
import zipfile
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

TEMPLATE_PATH = Path(__file__).resolve().parent.parent.parent / "xlsx_templates" / "ТКП.xlsx"

# Ячейки шаблона
CELL_NUMBER = "B6"  # «{№ ТКП} от {дата}»
CELL_INCOMING = "B7"  # вх. №
CELL_COMPANY = "C6"  # заказчик
CELL_POSITION = "C7"  # должность руководителя заказчика
CELL_DIRECTOR = "C8"  # руководитель заказчика
ROW_ITEM = 14  # образец строки блока
ROW_HEADER = 15  # образец заголовка раздела
ROW_TOTAL = 16  # Итого (далее НДС, Всего)
SAMPLE_ROWS = 2
# Строки ниже таблицы (до сдвига)
ROW_PAYMENT = 20  # «Условия оплаты: …\nСрок изготовления …»
ROW_EXECUTOR = 28  # Исп.: (B)
ROW_PHONE = 29  # Тел.: (B)
COLS = 5
# Высота строки блока: на строку текста; ширина столбца B — около 60 символов шрифта 10
LINE_HEIGHT = 12.75
LINE_CHARS = 60

PAYMENT_PREFIX = "Условия оплаты:"


@dataclass
class TkpRow:
    text: str
    quantity: int | None = None  # None — заголовок раздела
    price: float = 0.0


@dataclass
class TkpData:
    number: str
    incoming: str = ""
    company: str = ""
    position: str = ""
    director: str = ""
    rows: list[TkpRow] = field(default_factory=list)
    nds_percent: float = 20.0
    payment: str = ""  # «70% предоплата, 30% доплата-1»; пусто — как в шаблоне
    executor: str = ""
    phone: str = ""


def shift_rows(ws, idx: int, amount: int) -> None:
    """Вставляет (amount > 0) или удаляет (amount < 0) строки с idx, сдвигая объединения, высоты и область печати."""
    if not amount:
        return
    deleted_end = idx - amount if amount < 0 else idx  # удаляются строки idx … deleted_end - 1
    merges = [mr for mr in list(ws.merged_cells.ranges) if mr.max_row >= idx]
    for mr in merges:
        ws.unmerge_cells(str(mr))
    heights = {r: d.height for r, d in ws.row_dimensions.items() if r >= idx and d.height is not None}
    for r in [r for r in ws.row_dimensions if r >= idx]:
        del ws.row_dimensions[r]
    if amount > 0:
        ws.insert_rows(idx, amount)
    else:
        ws.delete_rows(idx, -amount)
    for r, h in heights.items():
        if r >= deleted_end:
            ws.row_dimensions[r + amount].height = h
    for mr in merges:
        if amount < 0 and mr.min_row < deleted_end and mr.max_row >= idx:
            continue  # объединение задевает удалённые строки
        if mr.min_row >= deleted_end:
            mr.shift(0, amount)
        ws.merge_cells(mr.coord)
    if ws.print_area:
        m = re.match(r"^(?:'?[^!]*'?!)?\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)$", ws.print_area.split(",")[0])
        if m and int(m.group(4)) >= idx:
            ws.print_area = f"{m.group(1)}{m.group(2)}:{m.group(3)}{int(m.group(4)) + amount}"


def _row_height(text: str) -> float:
    lines = sum(max(1, math.ceil(len(line) / LINE_CHARS)) for line in text.split("\n"))
    return max(15.0, lines * LINE_HEIGHT + 2)


def _payment_text(template_text: str, payment: str) -> str:
    """Первая строка условий — из «Порядка оплаты» ТКП, остальные (срок изготовления) — из шаблона."""
    lines = (template_text or "").split("\n")
    first = f"{PAYMENT_PREFIX} {payment}"
    if lines and lines[0].startswith(PAYMENT_PREFIX):
        lines[0] = first
    else:
        lines.insert(0, first)
    return "\n".join(lines)


def build_tkp_xlsx(data: TkpData, template: bytes | None = None) -> bytes:
    template = template if template is not None else TEMPLATE_PATH.read_bytes()
    wb = load_workbook(io.BytesIO(template))
    ws = wb.active

    ws[CELL_NUMBER] = data.number
    ws[CELL_INCOMING] = data.incoming or None
    ws[CELL_COMPANY] = data.company or None
    ws[CELL_POSITION] = data.position or None
    ws[CELL_DIRECTOR] = data.director.upper() if data.director else None

    item_style = [copy(ws.cell(ROW_ITEM, c)._style) for c in range(1, COLS + 1)]
    header_style = [copy(ws.cell(ROW_HEADER, c)._style) for c in range(1, COLS + 1)]
    header_height = ws.row_dimensions[ROW_HEADER].height
    n = len(data.rows)
    shift_rows(ws, ROW_ITEM + min(n, SAMPLE_ROWS), n - SAMPLE_ROWS)
    delta = n - SAMPLE_ROWS

    for i, row in enumerate(data.rows):
        r = ROW_ITEM + i
        styles = header_style if row.quantity is None else item_style
        for c in range(1, COLS + 1):
            ws.cell(r, c)._style = copy(styles[c - 1])
            ws.cell(r, c).value = None
        ws.cell(r, 2).value = row.text
        if row.quantity is None:
            ws.row_dimensions[r].height = header_height
            continue
        ws.cell(r, 3).value = row.quantity
        ws.cell(r, 4).value = round(row.price, 2)
        ws.cell(r, 5).value = f"=D{r}*C{r}"
        ws.row_dimensions[r].height = _row_height(row.text)

    total = ROW_TOTAL + delta
    ws.cell(total, 5).value = f"=SUM(E{ROW_ITEM}:E{total - 1})" if n else 0
    ws.cell(total + 1, 4).value = f"НДС {data.nds_percent:g}%:"
    ws.cell(total + 1, 5).value = f"=E{total}*{data.nds_percent / 100:g}"
    ws.cell(total + 2, 5).value = f"=E{total}+E{total + 1}"

    if data.payment:
        cell = ws.cell(ROW_PAYMENT + delta, 2)
        cell.value = _payment_text(cell.value, data.payment)
    if data.executor:
        ws.cell(ROW_EXECUTOR + delta, 2).value = data.executor
    if data.phone:
        ws.cell(ROW_PHONE + delta, 2).value = data.phone

    buf = io.BytesIO()
    wb.save(buf)
    return copy_drawing(template, buf.getvalue())


# ── Перенос логотипа (drawing первого листа) из шаблона ──

_RELS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_DRAWING_TYPE = f"{_R_NS}/drawing"
_DRAWING_REL_ID = "rIdTkpLogo"


def _sheet_drawing(zf: zipfile.ZipFile) -> str | None:
    """Путь drawing первого листа шаблона."""
    rels = "xl/worksheets/_rels/sheet1.xml.rels"
    if rels not in zf.namelist():
        return None
    for m in re.finditer(r"<Relationship\b[^>]*>", zf.read(rels).decode("utf-8")):
        tag = m.group(0)
        if f'Type="{_DRAWING_TYPE}"' in tag:
            target = re.search(r'Target="([^"]+)"', tag).group(1)
            return "xl/" + target.replace("../", "") if target.startswith("../") else target.lstrip("/")
    return None


def copy_drawing(template: bytes, result: bytes) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(template))
    drawing = _sheet_drawing(src)
    if not drawing:
        return result
    name = drawing.rsplit("/", 1)[1]
    drawing_rels = f"xl/drawings/_rels/{name}.rels"
    parts = {drawing: src.read(drawing)}
    if drawing_rels in src.namelist():
        rels_xml = src.read(drawing_rels)
        parts[drawing_rels] = rels_xml
        for target in re.findall(r'Target="\.\./([^"]+)"', rels_xml.decode("utf-8")):
            parts[f"xl/{target}"] = src.read(f"xl/{target}")

    dst = zipfile.ZipFile(io.BytesIO(result))
    out = io.BytesIO()
    sheet_rels = "xl/worksheets/_rels/sheet1.xml.rels"
    rel = f'<Relationship Id="{_DRAWING_REL_ID}" Type="{_DRAWING_TYPE}" Target="../drawings/{name}"/>'
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in dst.infolist():
            data = dst.read(item.filename)
            if item.filename in parts:
                continue
            if item.filename == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                tag = f'<drawing xmlns:r="{_R_NS}" r:id="{_DRAWING_REL_ID}"/>'
                m = re.search(
                    r"<(?:legacyDrawing|legacyDrawingHF|picture|oleObjects|controls|tableParts|extLst)\b", text
                )
                pos = m.start() if m else text.rindex("</worksheet>")
                data = (text[:pos] + tag + text[pos:]).encode("utf-8")
            elif item.filename == sheet_rels:
                data = data.decode("utf-8").replace("</Relationships>", rel + "</Relationships>").encode("utf-8")
            elif item.filename == "[Content_Types].xml":
                text = data.decode("utf-8")
                for ext in {p.rsplit(".", 1)[1].lower() for p in parts if p.startswith("xl/media/")}:
                    if f'Extension="{ext}"' not in text:
                        mime = "image/jpeg" if ext in ("jpg", "jpeg") else f"image/{ext}"
                        text = text.replace(
                            "<Default ", f'<Default Extension="{ext}" ContentType="{mime}"/><Default ', 1
                        )
                override = (
                    f'<Override PartName="/{drawing}" '
                    'ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>'
                )
                data = text.replace("</Types>", override + "</Types>").encode("utf-8")
            zout.writestr(item, data)
        if sheet_rels not in dst.namelist():
            zout.writestr(
                sheet_rels,
                f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="{_RELS_NS}">{rel}</Relationships>',
            )
        for path, data in parts.items():
            zout.writestr(path, data)
    return out.getvalue()
