"""Общие помощники для импорта/экспорта Excel."""

import io
from urllib.parse import quote

from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin")
)
# Разделитель групп — по локали Excel (в русской — пробел), поэтому «,» а не литеральный пробел:
# «# ##0.00» отделяло только последнюю тройку (1234 567,00)
MONEY_FMT = "#,##0.00"


def normalize(value) -> str:
    return str(value).strip().lower().replace("ё", "е") if value is not None else ""


def find_header(rows: list, aliases: dict[str, str], required: str, scan: int = 10) -> tuple[int, dict[str, int]]:
    """Ищет строку заголовка среди первых `scan` строк. Возвращает (индекс строки, {поле: индекс колонки})."""
    for row_idx, row in enumerate(rows[:scan]):
        mapping: dict[str, int] = {}
        for col_idx, cell in enumerate(row or ()):
            field = aliases.get(normalize(cell))
            if field and field not in mapping:
                mapping[field] = col_idx
        if required in mapping:
            return row_idx, mapping
    return -1, {}


def cell(row, mapping: dict[str, int], field: str):
    idx = mapping.get(field)
    if idx is None or idx >= len(row):
        return None
    value = row[idx]
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def to_float(value, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    try:
        return float(str(value).replace(" ", "").replace("\xa0", "").replace(",", "."))
    except (TypeError, ValueError):
        return default


def to_int(value, default: int | None = None) -> int | None:
    if value is None or value == "":
        return default
    try:
        return int(to_float(value, float(default) if default is not None else 0.0))
    except (TypeError, ValueError):
        return default


def data_rows(rows: list, header_idx: int):
    """Строки данных после заголовка — до первой полностью пустой строки."""
    for row in rows[header_idx + 1 :]:
        if not row or all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
            break
        yield row


def write_header(ws, headers: list[str], row_num: int | None = None) -> None:
    ws.append(headers)
    row_num = row_num or ws.max_row
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row_num, col_idx)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = THIN_BORDER


def border_row(ws, row_num: int, cols: int) -> None:
    for col_idx in range(1, cols + 1):
        ws.cell(row_num, col_idx).border = THIN_BORDER


def auto_fit_columns(ws) -> None:
    for col_cells in ws.columns:
        max_len = max((len(str(c.value)) for c in col_cells if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max_len + 3, 60)


def xlsx_response(wb: Workbook, filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


def safe_filename(name: str) -> str:
    return name.replace("/", "_").replace("\\", "_").replace(" ", "_")[:50]
